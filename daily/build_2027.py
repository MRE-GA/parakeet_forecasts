"""2027 daily sales for payroll: the chosen daily model, retrained on all data, pinned to the monthly figures sent.

Daily model (best in the daily backtests): the geometric mean of Prophet with the adopted settings and the extended
special days (daily/prophet_variants.py) and the pooled tree model (LightGBM, half-life 180 days), both retrained on
every day of data through Oct 14, 2026. It only sets
how each month divides into its days: each month is then split in proportion to it, in cents, so every month (and
year) matches what was sent exactly:
  * targets: 'Targets 2027' in outputs/Parakeet_2027_Manager_Targets.xlsx
  * run-rate: the run-rate objective by month behind the deck (report/build_targets.py), adding up to the annual
    figure in the workbook's 'Reference' tab

    python daily/build_2027.py   -> daily/outputs/Parakeet_2027_Daily_Forecast.xlsx, daily/outputs/daily_2027.csv
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from joblib import Parallel, delayed

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "report"))
from common import LAST_DAY, LGBM, NAME, OUT, ROOT, STORES, build_panel, future_panel, split_cents  # noqa: E402
from forecasting import load_stores  # noqa: E402
from prophet_variants import fit as prophet_fit  # noqa: E402

YEAR = pd.date_range("2027-01-01", "2027-12-31")
M27 = pd.period_range("2027-01", "2027-12", freq="M")
# stores whose within-month shape is pulled toward equal days (share of the way, in logs): Beverly Hills' weekly
# pattern has been flattening every year (weekends +30% in 2024, +20-25% in 2025, about 0% in 2026), so its learned
# shape overshoots; halfway to equal days cut its backtest error from 12.0% to 10.3% (10.7% to 10.1% in 2025, 12.9% to
# 10.5% in 2026) and on its holidays from 22% to 16%
SHRINK_TO_EQUAL = {"Beverly Hills": 0.5}
MONTHLY_XLSX = ROOT / "outputs" / "Parakeet_2027_Manager_Targets.xlsx"
CURRENT_DAILY = ROOT / "outputs" / "Parakeet_2027_Daily_Targets.xlsx"
XLSX = OUT / "Parakeet_2027_Daily_Forecast.xlsx"
CSV = OUT / "daily_2027.csv"
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
RECALC = Path.home() / ".claude" / "skills"
DISPLAY = [NAME.get(s, s) for s in STORES]


def sent_monthly() -> tuple[dict, dict]:
    """The monthly targets (as stored in the workbook) and the run-rate objective by month, adding up to the annual
    run-rate objective in the workbook's Reference tab."""
    from openpyxl import load_workbook

    import build_targets as bt

    wb = load_workbook(MONTHLY_XLSX, data_only=True)
    target = {r[0]: np.array(r[1:13], dtype=float) for r in wb["Targets 2027"].iter_rows(min_row=5, values_only=True)
              if r[0] in DISPLAY}
    ref = {r[0]: float(r[5]) for r in wb["Reference"].iter_rows(min_row=5, values_only=True) if r[0] in DISPLAY}
    n = bt.collect()
    run = {}
    for s, name in zip(STORES, DISPLAY):
        path = np.asarray(n["C"][s]["path"][3:], dtype=float)
        assert abs(path.sum() - ref[name]) < 0.02, (name, path.sum(), ref[name])
        run[name] = split_cents(ref[name], path) / 100  # months in cents, adding up to the annual figure sent
    return target, run


def daily_weights(panel: pd.DataFrame, stores: dict) -> dict[str, np.ndarray]:
    origin = pd.Timestamp(LAST_DAY)
    prophet = Parallel(n_jobs=4)(delayed(prophet_fit)(stores[s], origin, YEAR, "tuned", "ext") for s in STORES)
    tree = LGBM(panel, origin, half_life=180)
    out = {}
    for s, name, wp in zip(STORES, DISPLAY, prophet):
        fp = future_panel(s, YEAR, stores[s].open_flags(YEAR))
        w = np.sqrt(wp * tree.weights(fp))
        a = SHRINK_TO_EQUAL.get(s, 0.0)
        out[name] = w ** (1 - a) * (w > 0)  # equal days have weight 1 on every open day
    return out


def allocate(monthly: dict, weights: dict) -> dict[str, np.ndarray]:
    months = YEAR.month
    cents = {}
    for name in DISPLAY:
        c = np.concatenate([split_cents(monthly[name][i], weights[name][months == i + 1]) for i in range(12)])
        for i in range(12):
            assert c[months == i + 1].sum() == int(round(monthly[name][i] * 100)), (name, i)
        assert (c[weights[name] == 0] == 0).all()
        cents[name] = c
    return cents


def export(t_cents: dict, r_cents: dict, target: dict, run: dict) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    arial = "Arial"
    title, bold, normal = (Font(name=arial, size=14, bold=True), Font(name=arial, size=10, bold=True),
                           Font(name=arial, size=10))
    blue, green = Font(name=arial, size=10, color="0000FF"), Font(name=arial, size=10, color="008000")
    muted = Font(name=arial, size=9, italic=True, color="565A5C")
    head_fill, thin = PatternFill("solid", fgColor="F2F2F2"), Side(style="thin", color="000000")
    money = '$#,##0.00;($#,##0.00);"-"'

    def header(ws, row, labels):
        for j, h in enumerate(labels, start=1):
            c = ws.cell(row, j, h)
            c.font, c.fill, c.border = bold, head_fill, Border(bottom=thin)
            c.alignment = Alignment(horizontal="left" if j == 1 else "center", wrap_text=True)

    wb = Workbook()
    first, last = 5, 4 + len(YEAR)
    sheets = {}
    for k, (sheet, cents, what) in enumerate((("Daily targets 2027", t_cents, "the monthly manager targets"),
                                             ("Daily run-rate 2027", r_cents, "the run-rate objective by month"))):
        ws = wb.active if k == 0 else wb.create_sheet(sheet)
        ws.title = sheet
        sheets[sheet] = ws
        ws["A1"], ws["A1"].font = f"Parakeet Café · 2027 daily sales: {what}", title
        ws["A2"], ws["A2"].font = (
            f"Each month of {what} (as sent) split over its days by the daily model: Prophet (with every special day) "
            "and a tree model trained on all stores, blended, retrained on data through Oct 14, 2026. Every month adds "
            "up exactly to the figure sent (see 'Check'). Blue = values from the analysis; black = formulas."), muted
        header(ws, 4, ["Date", "Day", "Month"] + DISPLAY + ["Total"])
        for i, day in enumerate(YEAR):
            r = first + i
            c = ws.cell(r, 1, day.to_pydatetime())
            c.number_format, c.font = "yyyy-mm-dd", normal
            ws.cell(r, 2, DAYS[day.dayofweek]).font = normal
            ws.cell(r, 3, day.strftime("%b %Y")).font = normal
            for j, name in enumerate(DISPLAY):
                c = ws.cell(r, 4 + j, cents[name][i] / 100)
                c.font, c.number_format = blue, money
            c = ws.cell(r, 4 + len(DISPLAY), f"=SUM(D{r}:{get_column_letter(3 + len(DISPLAY))}{r})")
            c.font, c.number_format = normal, money
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 12, 6, 10
        for j in range(4, 5 + len(DISPLAY)):
            ws.column_dimensions[get_column_letter(j)].width = 14
        ws.row_dimensions[4].height = 30
        ws.freeze_panes = "D5"

    # ---- weekday averages by season: a staffing template ----------------------------------------------------
    wa = wb.create_sheet("Weekday averages")
    wa["A1"], wa["A1"].font = "Average daily sales by weekday and season, 2027 (holidays excluded)", title
    wa["A2"], wa["A2"].font = ("Seasons: summer Jun 15 – Aug 20; holiday season Nov 22 – Jan 5; spring break the week "
                               "either side of Easter; school term the rest. From the two daily tabs."), muted
    header(wa, 4, ["Store", "Series", "Season"] + DAYS)
    from common import SPECIAL, season_of

    seasons = season_of(YEAR)
    plain = ~YEAR.isin(SPECIAL.index)
    r = 5
    for name in DISPLAY:
        for series, cents in (("Target", t_cents), ("Run-rate", r_cents)):
            for s in ["term", "summer", "holidays"]:
                mask = (seasons == s) & plain
                v = pd.Series(cents[name][mask] / 100, index=YEAR[mask]).groupby(YEAR[mask].dayofweek).mean()
                wa.cell(r, 1, name).font = normal
                wa.cell(r, 2, series).font = normal
                wa.cell(r, 3, {"term": "School term", "summer": "Summer", "holidays": "Holiday season"}[s]).font = normal
                for k in range(7):
                    c = wa.cell(r, 4 + k, round(float(v.get(k, 0.0)), 2))
                    c.font, c.number_format = blue, '$#,##0'
                r += 1
    wa.column_dimensions["A"].width, wa.column_dimensions["B"].width, wa.column_dimensions["C"].width = 20, 10, 16
    for k in range(7):
        wa.column_dimensions[get_column_letter(4 + k)].width = 10
    wa.freeze_panes = "D5"

    # ---- check against the monthly figures sent ---------------------------------------------------------------
    ck = wb.create_sheet("Check")
    ck["A1"], ck["A1"].font = "Check: the daily figures add up to the monthly figures sent", title
    ck["A2"], ck["A2"].font = ("Monthly targets: 'Targets 2027' in Parakeet_2027_Manager_Targets.xlsx. Run-rate: the "
                               "run-rate objective by month, adding up to 'Reference'. Difference in cents: must be 0."
                               ), muted
    ck["A3"], ck["A3"].font = "Total absolute difference (cents)", bold
    header(ck, 5, ["Store", "Series", "Month", "Monthly (as sent)", "Sum of daily", "Difference (cents)"])
    r = 6
    for series, sheet, monthly in (("Target", "Daily targets 2027", target), ("Run-rate", "Daily run-rate 2027", run)):
        for j, name in enumerate(DISPLAY):
            col = get_column_letter(4 + j)
            start = r
            for i, label in enumerate(M27.strftime("%b %Y")):
                ck.cell(r, 1, name).font = normal
                ck.cell(r, 2, series).font = normal
                ck.cell(r, 3, label).font = normal
                c = ck.cell(r, 4, float(monthly[name][i]))
                c.font, c.number_format = blue, money
                c = ck.cell(r, 5, f"=SUMIFS('{sheet}'!{col}${first}:{col}${last},'{sheet}'!$C${first}:$C${last},C{r})")
                c.font, c.number_format = green, money
                ck.cell(r, 6, f"=ROUND(E{r}*100,0)-ROUND(D{r}*100,0)").font = normal
                r += 1
            ck.cell(r, 1, name).font = bold
            ck.cell(r, 2, series).font = bold
            ck.cell(r, 3, "Year").font = bold
            for k, L in ((4, "D"), (5, "E")):
                c = ck.cell(r, k, f"=SUM({L}{start}:{L}{r - 1})")
                c.font, c.number_format = bold, money
            ck.cell(r, 6, f"=ROUND(E{r}*100,0)-ROUND(D{r}*100,0)").font = bold
            for k in range(1, 7):
                ck.cell(r, k).border = Border(top=thin)
            r += 2
    ck["B3"] = f"=SUMPRODUCT(ABS(F6:F{r - 2}))"
    ck["B3"].font = bold
    for L, wdt in (("A", 20), ("B", 12), ("C", 12), ("D", 20), ("E", 20), ("F", 18)):
        ck.column_dimensions[L].width = wdt
    ck.freeze_panes = "A6"

    # ---- notes -------------------------------------------------------------------------------------------------
    nt = wb.create_sheet("Notes")
    nt["A1"], nt["A1"].font = "How the daily figures are built", title
    for i, text in enumerate([
        "Monthly and annual figures: exactly those sent (targets workbook and deck); this file only splits each month "
        "into its days.",
        "Daily model: the geometric mean of two models, both retrained on all data through Oct 14, 2026: Prophet with "
        "the adopted settings (time of year, day of week, holidays) plus the other special days (Dec 26, the days "
        "between Christmas and New Year's Eve, Thanksgiving weekend, spring break, Halloween), and a tree model "
        "(LightGBM) trained on all ten stores together on each day's sales against its month's average, from the "
        "weekday, season, special days and time of year, with recent days weighing more.",
        "Chosen by backtest: six forecast origins (Dec 2024 to Mar 2026), each month of the next year split into "
        "days and compared with actual days, given the true month. Daily error: Prophet as tuned 9.2%, this blend "
        "8.5%; on days-off holidays 12.8% and 11.8%. Tested and not used: weekly patterns by season or for recent "
        "months in Prophet, seasonality within the month, holidays set to their usual level.",
        "Beverly Hills: its daily shape is pulled halfway toward equal days. Its weekly pattern has been flattening "
        "every year (weekends about +30% over weekdays in 2024, +20-25% in 2025, about 0% in 2026), so the learned "
        "shape overshoots; halfway cut its backtest error from 12.0% to 10.3%, in 2025 and in 2026.",
        "Each month is split in cents in proportion to the model's days; amounts are rounded so the month adds up "
        "exactly (see 'Check').",
        "Holidays the stores have always traded on are kept as trading days; a café closure that is not in the "
        "calendar will show up as a shortfall on that day.",
    ], start=3):
        nt.cell(i, 1, text).font = normal
    nt.column_dimensions["A"].width = 140

    wb.save(XLSX)
    script = next(RECALC.rglob("xlsx/scripts/recalc.py"), None)
    if script is None:
        return {"status": "not recalculated"}
    res = subprocess.run([sys.executable, str(script), str(XLSX), "180"], capture_output=True, text=True)
    out = json.loads(res.stdout[res.stdout.index("{"):])
    if out.get("status") != "success":
        raise RuntimeError(f"workbook has formula errors: {out}")
    return out


def main() -> None:
    target, run = sent_monthly()
    panel = build_panel()
    stores = load_stores(LAST_DAY)
    w = daily_weights(panel, stores)
    t_cents, r_cents = allocate(target, w), allocate(run, w)
    rows = [{"date": d.date(), "weekday": DAYS[d.dayofweek], "store": name, "target": t_cents[name][i] / 100,
             "runrate": r_cents[name][i] / 100} for name in DISPLAY for i, d in enumerate(YEAR)]
    pd.DataFrame(rows).to_csv(CSV, index=False, float_format="%.2f")
    print("workbook:", export(t_cents, r_cents, target, run))
    from openpyxl import load_workbook

    diff = load_workbook(XLSX, data_only=True)["Check"]["B3"].value
    assert diff == 0, diff
    print(f"check: total absolute difference {diff} cents")
    for name in DISPLAY:
        print(f"{name:18s} target {t_cents[name].sum() / 100:>14,.2f} (sent {target[name].sum():>14,.2f})   "
              f"run-rate {r_cents[name].sum() / 100:>14,.2f} (sent {run[name].sum():>14,.2f})")
    print(f"-> {XLSX.relative_to(ROOT)}\n-> {CSV.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
