"""Import the investor budget workbook (2027_Sales_budget_Investors.xlsx) into data/.

    .venv/bin/python -I src/import_budget.py path/to/2027_Sales_budget_Investors.xlsx

Reads cached cell values only (openpyxl, data_only) and writes:
  data/cfo_budget_2027.csv              per café: 2025 actual, 2026 est. finish, 2027 objective, key assumption
  data/cfo_budget_2027_monthly.csv      per café and month: 2025 actual, 2026 actual (Jan-Sep sales report) or
                                        estimate (Oct-Dec), 2027 objective (incl. new and remodeled cafés)
  data/cfo_budget_2027_comp_trend.csv   the budget's same-store comp trend and number of comp cafés

Standalone on purpose (no project imports) so it can run in isolated mode on an untrusted file.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

DATA = Path(__file__).resolve().parents[1] / "data"
# budget name -> Toast store name (None: not trading yet)
STORE = {"Beverly Hills": "Beverly Hills", "Brentwood": "Brentwood", "Carlsbad": "Carlsbad", "Coronado": "Coronado",
         "Del Mar Café": "Carmel Valley", "La Jolla Café": "La Jolla", "La Jolla Juicery": "La Jolla Juicery",
         "Little Italy Café": "Little Italy", "Newport Beach": "Newport", "La Jolla Cove": "La Jolla Cove",
         "Manhattan Beach": None, "Sportsmen's Lodge": None, "Boulangerie": None, "South Coast Plaza": None,
         "Riverwalk": None}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def rows_by_label(ws, col: int = 2) -> dict[str, int]:
    """First row number for each text label in a column."""
    out = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip() and v.strip() not in out:
            out[v.strip()] = r
    return out


def split_name(label: str) -> tuple[str, str | None, str | None]:
    """'Manhattan Beach (opens Mar 1)' -> ('Manhattan Beach', 'new', '2027-03-01')."""
    m = re.match(r"^(.*?)\s*\((.*)\)$", label)
    if not m:
        return label, None, None
    name, note = m.group(1).strip(), m.group(2)
    if "2028" in note:
        return name, "new", "2028"
    d = re.search(r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w* (\d{1,2})", note)
    opens = f"2027-{MONTHS.index(d.group(1)) + 1:02d}-{int(d.group(2)):02d}" if d else None
    return name, ("remodel" if "remodel" in note else "new"), opens


def number(v):
    return float(v) if isinstance(v, (int, float)) else None


def main(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)

    # ---- per café (Summary) ------------------------------------------------------------
    ws = wb["Summary"]
    header = rows_by_label(ws)["Café"]
    assert [ws.cell(header, c).value for c in range(2, 10)] == [
        "Café", "2025 Actual", "2026 Est. Finish", "2026 Growth", "2027 Objective", "2027 Growth", "$ Increase",
        "Key Assumption"], "Summary layout changed"
    cafes, section = [], "existing"
    for r in range(header + 1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if not isinstance(label, str) or not label.strip():
            continue
        label = label.strip()
        if label.startswith("NEW & REMODELED CAFÉS"):
            section = "new"
            continue
        if label.endswith("Total") or label.startswith("TOTAL") or label.startswith("2026 Est.") or \
                label.startswith("2027 objectives") or label.startswith("La Jolla Cove reopens"):
            continue
        name, status, opens = split_name(label)
        if name not in STORE:
            raise ValueError(f"unknown café in the budget: {label!r}")
        objective = number(ws.cell(r, 6).value)
        if opens == "2028" or (status == "new" and objective is None):
            objective = None  # no 2027 volume (opens 2028, or volume TBD)
        cafes.append({
            "budget_name": name, "store": STORE[name], "status": status or section, "opens": opens,
            "actual_2025": number(ws.cell(r, 3).value), "est_2026": number(ws.cell(r, 4).value),
            "objective_2027": objective,
            "growth_2027": number(ws.cell(r, 7).value) if status != "new" else None,
            "increase_2027": number(ws.cell(r, 8).value) if objective is not None else None,
            "key_assumption": ws.cell(r, 9).value,
        })
    per_cafe = pd.DataFrame(cafes)

    # ---- monthly: 2025 actual and 2026 actual/estimate (2026 Base), 2027 objective (2027 Monthly) -----
    rows = []
    base = wb["2026 Base"]
    labels = rows_by_label(base)
    for title, year in (("2025 Actual", 2025), ("2026 Actual / Est.", 2026)):
        r0 = labels[title]
        assert [base.cell(r0, c).value for c in range(3, 15)] == MONTHS, f"{title} layout changed"
        r = r0 + 1
        while isinstance(base.cell(r, 2).value, str) and base.cell(r, 2).value.strip() in STORE:
            name = base.cell(r, 2).value.strip()
            for i in range(12):
                basis = "actual" if year == 2025 or i < 9 else "estimate"
                rows.append({"budget_name": name, "store": STORE[name], "month": f"{year}-{i + 1:02d}",
                             "basis": basis, "value": number(base.cell(r, 3 + i).value)})
            r += 1
    m27 = wb["2027 Monthly"]
    r0 = rows_by_label(m27)["Café"]
    assert [m27.cell(r0, c).value for c in range(3, 15)] == MONTHS, "2027 Monthly layout changed"
    for r in range(r0 + 1, m27.max_row + 1):
        label = m27.cell(r, 2).value
        if not isinstance(label, str):
            continue
        name = split_name(label.strip())[0]
        if name not in STORE or (per_cafe.loc[per_cafe.budget_name == name, "objective_2027"].isna().all()):
            continue
        for i in range(12):
            rows.append({"budget_name": name, "store": STORE[name], "month": f"2027-{i + 1:02d}", "basis": "objective",
                         "value": number(m27.cell(r, 3 + i).value) or 0.0})
    monthly = pd.DataFrame(rows)

    # ---- comp trend (Comp History) --------------------------------------------------------
    ch = wb["Comp History"]
    lab = rows_by_label(ch)
    hdr = lab["Café"]
    cols = {ch.cell(hdr, c).value: c for c in range(2, ch.max_column + 1)}
    trend = pd.DataFrame([
        {"period": p, "same_store_growth": number(ch.cell(lab["Same-Store Comp Total"], cols[h]).value),
         "comp_cafes": int(ch.cell(lab["# of comp cafés"], cols[h]).value)}
        for p, h in (("2024", "2024 Comp"), ("2025", "2025 Comp"), ("2026E", "2026 Est. Comp"),
                     ("2027 budget", "2027 Plan"))])

    # ---- checks: months add up to the annual figures --------------------------------------
    for _, c in per_cafe.iterrows():
        mm = monthly[monthly.budget_name == c.budget_name]
        if c.est_2026 is not None and not pd.isna(c.est_2026):
            got = mm[mm.month.str.startswith("2026")].value.sum()
            assert abs(got - c.est_2026) < 0.01, (c.budget_name, got, c.est_2026)
        if c.objective_2027 is not None and not pd.isna(c.objective_2027):
            got = mm[mm.basis == "objective"].value.sum()
            assert abs(got - c.objective_2027) < 0.01, (c.budget_name, got, c.objective_2027)

    per_cafe.to_csv(DATA / "cfo_budget_2027.csv", index=False, float_format="%.2f")
    monthly.to_csv(DATA / "cfo_budget_2027_monthly.csv", index=False, float_format="%.2f")
    trend.to_csv(DATA / "cfo_budget_2027_comp_trend.csv", index=False, float_format="%.6f")
    ex = per_cafe[per_cafe.status == "existing"]
    print(per_cafe.to_string())
    print(f"existing: 2026E {ex.est_2026.sum():,.2f}  2027 {ex.objective_2027.sum():,.2f}  "
          f"total company 2027 {per_cafe.objective_2027.sum():,.2f}")
    print(trend.to_string())


if __name__ == "__main__":
    main(sys.argv[1])
