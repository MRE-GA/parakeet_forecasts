"""Build the week-types deck for the next 52 weeks (payroll templates: sales and checks by hour and day for each type
of week, store by store) and its workbook.

    python report/build_week_types.py [--ds PATH_TO_PARAKEET_DESIGN_SYSTEM] [--no-pdf] [--shots DIR]

The weeks, types and templates come from weektypes/ (model.py, hourly.py, templates.py; backtests in
backtest_hourly.py), on today's run-rate by day (run report/build_daily_targets.py first) and the hourly Toast extract
(data/raw/parakeet_toast_hourly.csv.gz); the live check from weektypes/live_check.py. Outputs:
outputs/Parakeet_Week_Types_Next_52_Weeks.pdf and outputs/Parakeet_Week_Templates_Next_52_Weeks.xlsx.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from jinja2 import Environment, FileSystemLoader

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "weektypes"))
sys.path.insert(0, str(ROOT / "report"))

from build_report import BUILD, copy_assets, find_design_system  # noqa: E402  (first: it imports src/backtest.py)
import templates as T  # noqa: E402
from hourly import CHECK_MONTHS, DAYPARTS, RECENCY_HALF_LIFE, SEASON_SHIFT  # noqa: E402
from model import (COVE, DAYS, DISPLAY, HOURS_FROM, LAST_DAY, MIN_RUN, N_WEEKS, REMODELED, STEP, STORES,  # noqa: E402
                   runrate_days, week_start)

PDF = ROOT / "outputs" / "Parakeet_Week_Types_Next_52_Weeks.pdf"
XLSX = ROOT / "outputs" / "Parakeet_Week_Templates_Next_52_Weeks.xlsx"
RECALC = Path.home() / ".claude" / "skills"
BACKTEST = ROOT / "outputs" / "weektypes" / "backtest_hourly.csv"
LIVE = ROOT / "outputs" / "weektypes" / "live_check.csv"
COLOR = {"Low": "#cccccc", "Normal": "#82888a", "High": "#f59ab8", "Super high": "#ef4b81", REMODELED: "#373a3c"}
INK_ON = {"Low": "#000", "Normal": "#fff", "High": "#000", "Super high": "#fff", REMODELED: "#fff"}
PINK = (239, 75, 129)
NUM = {2: "two", 3: "three", 4: "four", 5: "five"}
SHORT = {"Fourth of July week": "Fourth of July", "Week after the Fourth": "After the Fourth",
         "Thanksgiving week": "Thanksgiving", "Christmas week": "Christmas", "New Year's week": "New Year's",
         "Memorial Day week": "Memorial Day", "Father's Day week": "Father's Day", "Spring break": "Spring break",
         "Veterans Day week": "Veterans Day", "Reopening week": "Reopening"}


# ------------------------------------------------------------------------------------------
# Formatting
# ------------------------------------------------------------------------------------------
def money(x: float) -> str:
    return f"${x:,.0f}"


def fk(x: float) -> str:
    return f"${x / 1000:,.1f}K"


def pct(x: float) -> str:
    return f"{x * 100:+.0f}%".replace("-", "−")


def pct1(x: float) -> str:
    return f"{x * 100:+.1f}%".replace("-", "−")


def fm(x: float) -> str:
    return f"${x / 1e6:,.2f}M"


def dspan(a: pd.Timestamp, b: pd.Timestamp) -> str:
    """Nov 23–29, 2026; Mar 8 – Jun 27, 2027; Dec 28, 2026 – Jan 3, 2027."""
    if a.year != b.year:
        return f"{a:%b} {a.day}, {a.year} – {b:%b} {b.day}, {b.year}"
    if a.month == b.month:
        return f"{a:%b} {a.day}–{b.day}, {b.year}"
    return f"{a:%b} {a.day} – {b:%b} {b.day}, {b.year}"


def week_range(a: int, b: int) -> dict:
    first, last = week_start(a), week_start(b) + pd.Timedelta(days=6)
    label = f"Weeks {a}–{b}" if a != b else f"Week {a}"
    return {"weeks": label, "dates": dspan(first, last), "n": b - a + 1}


def mlabel(d: pd.Timestamp) -> str:
    """A month on an axis; January with its year."""
    return f"{d:%b} {d.year}" if d.month == 1 else f"{d:%b}"


def and_join(items: list[str]) -> str:
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


# ------------------------------------------------------------------------------------------
# Pieces of the pages
# ------------------------------------------------------------------------------------------
def shade(v: float, vmax: float) -> tuple[str, str]:
    t = max(0.0, min(1.0, v / vmax)) ** 1.1 if vmax > 0 else 0.0
    r, g, b = (round(255 + (c - 255) * t) for c in PINK)
    return f"rgb({r},{g},{b})", ("#fff" if t > 0.58 else "#000")


def hour_label(h: int) -> str:
    return f"{(h - 1) % 12 + 1}{'am' if h < 12 else 'pm'}"


def day_parts(t: T.Template) -> list[float]:
    """Share of the week's sales before 11am, from 11am to 2pm and after 2pm."""
    tot = t.sales.sum()
    return [t.sales[[i for i, hr in enumerate(t.hours) if lo <= hr <= hi]].sum() / tot for lo, hi in DAYPARTS]


def heat_table(t: T.Template, vmax: float) -> str:
    special = [lab.replace(" (observed)", ", observed") for lab in (t.special or [""] * 7)]
    head = "".join(f"<th class='sp'>{c}<small>{lab}</small></th>" if lab else f"<th>{c}</th>"
                   for c, lab in zip(t.columns, special))
    rows = [f"<table class='heat'><thead><tr><th class='hr'></th>{head}<th class='tot'>Week</th></tr></thead><tbody>"]
    for i, h in enumerate(t.hours):
        cells = []
        for j in range(7):
            bg, fg = shade(t.sales[i, j], vmax)
            cells.append(f"<td style='background:{bg};color:{fg}'><b>{money(t.sales[i, j])}</b>"
                         f"<span>{t.checks[i, j]:,} checks</span></td>")
        rows.append(f"<tr><th class='hr'>{hour_label(h)}</th>{''.join(cells)}"
                    f"<td class='tot'><b>{money(t.sales[i].sum())}</b>"
                    f"<span>{t.checks[i].sum():,} checks</span></td></tr>")
    rows.append("<tr class='day'><th class='hr'>Day</th>" + "".join(
        f"<td><b>{money(t.day_sales[j])}</b><span>{t.day_checks[j]:,} checks</span></td>" for j in range(7))
        + f"<td class='tot'><b>{money(t.week_sales)}</b><span>{t.checks.sum():,} checks</span></td></tr>")
    return "".join(rows) + "</tbody></table>"


def weekly_svg(sw, temps: list[T.Template], h: int, w: int = 1160) -> str:
    """The run-rate by week, coloured by type; holiday weeks with a page of their own outlined."""
    pad_l, pad_r, pad_t, pad_b = 54, 120, 14, 26
    y = sw.weekly.to_numpy()
    ymax = float(np.ceil(y.max() * 1.08 / 5000) * 5000)
    bw = (w - pad_l - pad_r) / N_WEEKS
    Y = lambda v: pad_t + (1 - v / ymax) * (h - pad_t - pad_b)  # noqa: E731
    s = [f'<svg viewBox="0 0 {w} {h}" width="{w}" height="{h}" class="wk">']
    step = 5000 if ymax <= 40000 else 10000 if ymax <= 80000 else 20000
    for v in np.arange(0, ymax + 1, step):
        s.append(f'<line x1="{pad_l}" x2="{w - pad_r}" y1="{Y(v):.1f}" y2="{Y(v):.1f}" stroke="#eceeef"/>'
                 f'<text x="{pad_l - 8}" y="{Y(v) + 4:.1f}" text-anchor="end" class="ax">${v / 1000:.0f}K</text>')
    months_done = set()
    for i in range(N_WEEKS):
        wk = i + 1
        t = sw.label[wk]
        x = pad_l + i * bw
        own = wk in sw.holiday
        s.append(f'<rect x="{x + 1.2:.1f}" y="{Y(y[i]):.1f}" width="{bw - 2.4:.1f}" height="{Y(0) - Y(y[i]):.1f}" '
                 f'fill="{COLOR[t]}"' + (' stroke="#000" stroke-width="1.6"' if own else "") + "/>")
        if own:
            s.append(f'<circle cx="{x + bw / 2:.1f}" cy="{Y(y[i]) - 7:.1f}" r="2.6" fill="#000"/>')
        m = (week_start(wk).year, week_start(wk).month)
        if m not in months_done and week_start(wk).day <= 7:
            months_done.add(m)
            s.append(f'<text x="{x + 1:.1f}" y="{h - 8}" class="ax">{mlabel(week_start(wk))}</text>')
    by = {tt.name: tt for tt in temps}
    for t in sw.types:
        v = by[t].week_sales
        col = COLOR[t] if t != "Low" else "#82888a"  # the Low grey is too light for a line
        s.append(f'<line x1="{pad_l}" x2="{w - pad_r + 6}" y1="{Y(v):.1f}" y2="{Y(v):.1f}" stroke="{col}" '
                 f'stroke-width="1.3" stroke-dasharray="4 3"/>')
    # type labels at the right, spread so they do not overlap
    ys = sorted(((Y(by[t].week_sales), t) for t in sw.types), reverse=True)
    placed = []
    for yy, t in ys:
        yy = min(yy, placed[-1] - 15) if placed else yy
        placed.append(yy)
        s.append(f'<text x="{w - pad_r + 10}" y="{yy + 4:.1f}" class="tl">{t} {fk(by[t].week_sales)}</text>')
    return "".join(s) + "</svg>"


def calendar_rows(weeks) -> list[dict]:
    rows = []
    for s in STORES:
        sw = weeks[s]
        cells = [{"color": COLOR[sw.label[w]], "own": w in sw.holiday,
                  "tip": f"Week {w}: {sw.holiday.get(w, sw.label[w])}"} for w in range(1, N_WEEKS + 1)]
        n_reg = len([t for t in sw.types if t != REMODELED])
        rows.append({"name": DISPLAY[s], "cells": cells, "n": n_reg, "remodeled": REMODELED in sw.types,
                     "n_hol": len(sw.holiday)})
    return rows


def month_heads() -> list[str]:
    """The calendar's month row: a month over its first full week, January with its year."""
    return [mlabel(week_start(w)) if week_start(w).day <= 7 else "" for w in range(1, N_WEEKS + 1)]


def week_heads() -> list[str]:
    """The calendar's second row: every fourth week's Monday."""
    return [f"{week_start(w):%b} {week_start(w).day}" if w % 4 == 1 else "" for w in range(1, N_WEEKS + 1)]


def chips(ty_weeks: list[int], wt: pd.DataFrame, own: dict[int, str]) -> list[dict]:
    """Every week a template is used in: its Monday and its own run-rate against the template."""
    out = []
    for w in ty_weeks:
        r = wt.loc[w]
        label = f"{week_start(w):%b} {week_start(w).day}"
        if w in own:
            out.append({"label": label, "vs": None, "own": SHORT.get(own[w], own[w])})
        else:
            out.append({"label": label, "vs": r.vs_template, "own": None})
    return out


# ------------------------------------------------------------------------------------------
# The numbers behind every page
# ------------------------------------------------------------------------------------------
def collect() -> dict:
    rr = runrate_days()
    weeks, temps = T.build(rr=rr)
    wt_all = T.week_table(weeks, temps, rr)
    stores = []
    for s in STORES:
        sw = weeks[s]
        by = {t.name: t for t in temps[s]}
        wt = wt_all[wt_all.store == DISPLAY[s]].set_index("week")
        reg_w = wt[wt.kind == "type"]
        vmax = max(t.sales.max() for t in temps[s] if t.kind == "type")
        regular_types = [t for t in sw.types if t != REMODELED]
        types = []
        for i, name in enumerate(sw.types):
            t = by[name]
            reg = sw.regular(name)
            below = sw.types[i - 1] if 0 < i and name != REMODELED else None
            ranges = [week_range(a, b) for a, b in sw.ranges(name)]
            d = t.day_sales
            pk = np.unravel_index(t.sales.argmax(), t.sales.shape)
            # holiday days in this type's weeks: holiday weeks without a page of their own
            notes = []
            for w in sorted(set(sw.label.index[sw.label == name]) & set(sw.holiday_days)):
                if w in sw.holiday:
                    continue
                for dt, lab in sw.holiday_days[w]:
                    if lab == "Spring break":
                        continue
                    v = rr.loc[dt, s]
                    notes.append({"day": lab.replace(" (observed)", ", observed"),
                                  "date": f"{dt:%a %b} {dt.day}, {dt.year}", "sales": v, "vs": v / d[dt.dayofweek] - 1})
            notes.sort(key=lambda r: -abs(r["vs"]))
            keep = 3 if name == REMODELED or len(sw.ranges(name)) > 1 else 4  # room left in the side column
            own = [{"name": sw.holiday[w], "week": w} for w in sorted(sw.holiday) if sw.label[w] == name]
            weekly = sw.weekly[reg]
            fit = reg_w[reg_w.template == name]
            types.append({
                "name": name, "color": COLOR[name], "ink": INK_ON[name], "t": t, "heat": heat_table(t, vmax),
                "n_weeks": len(t.weeks), "n_regular": len(reg), "ranges": ranges, "avg": t.week_sales,
                "below": below, "vs_below": (t.week_sales / by[below].week_sales - 1) if below else None,
                "checks": int(t.checks.sum()), "avg_check": t.week_sales / max(t.checks.sum(), 1),
                "peak": f"{DAYS[pk[1]]} {hour_label(t.hours[pk[0]])}", "peak_sales": t.sales[pk],
                "weekend": (d[5] + d[6]) / 2 / d[:5].mean() - 1, "lo": weekly.min(), "hi": weekly.max(),
                "busiest_day": DAYS[int(np.argmax(d))], "notes": notes[:keep], "more_notes": max(len(notes) - keep, 0),
                "more_max": max((abs(x["vs"]) for x in notes[keep:]), default=0.0),
                "own": own, "chips": chips(t.weeks, wt, sw.holiday), "parts": day_parts(t),
                "fit_week": fit.vs_template.abs().mean(), "fit_day": fit.day_gap_abs.sum() / fit.week_sales.sum(),
            })
        hols = []
        for w, name in sw.holiday.items():
            t = by[name]
            base = by[sw.label[w]]
            days = []
            for j, (dt, lab) in enumerate(zip(t.dates, t.special)):
                if lab:
                    days.append({"day": lab.replace(" (observed)", ", observed"),
                                 "date": f"{dt:%a %b} {dt.day}, {dt.year}", "sales": t.day_sales[j],
                                 "vs": t.day_sales[j] / base.day_sales[j] - 1})
            first, last = t.dates[0], t.dates[-1]
            hols.append({"name": name, "short": SHORT.get(name, name), "week": w, "t": t,
                         "heat": heat_table(t, vmax), "dates": dspan(first, last), "type": sw.label[w],
                         "type_color": COLOR[sw.label[w]], "type_ink": INK_ON[sw.label[w]], "total": t.week_sales,
                         "vs": t.week_sales / base.week_sales - 1, "base": base.week_sales, "days": days,
                         "checks": int(t.checks.sum()), "avg_check": t.week_sales / max(t.checks.sum(), 1),
                         "parts": day_parts(t)})
        steps = [ty["vs_below"] for ty in types if ty["vs_below"] is not None]
        rows = len(sw.types) + len(sw.holiday)
        small_step = len(regular_types) == 2 and steps and steps[0] < STEP
        stores.append({"key": s, "name": DISPLAY[s], "weeks": sw, "types": types, "hols": hols,
                       "n": len(regular_types), "n_word": NUM[len(regular_types)],
                       # the chart gives way to the table on stores with many rows
                       "svg": weekly_svg(sw, temps[s], h=max(130, min(260, 470 - 32 * rows))),
                       "small_step": small_step, "step": steps[0] if steps else None,
                       "hours": f"{hour_label(temps[s][0].hours[0])}–{hour_label(temps[s][0].hours[-1] + 1)}",
                       "cove": s == COVE, "fit_week": reg_w.vs_template.abs().mean(),
                       "fit_day": reg_w.day_gap_abs.sum() / reg_w.week_sales.sum(),
                       "far": reg_w.vs_template.abs().max()})
    bt = pd.read_csv(BACKTEST)
    hl = RECENCY_HALF_LIFE
    agg = bt.groupby(["test", "method"])[["err", "sales"]].sum()
    err = (agg.err / agg.sales * 100).round(2)
    live = pd.read_csv(LIVE)
    lv = live.groupby("store", sort=False)[["actual", "runrate", "hour_err", "checks", "checks_expected"]].sum()
    lv_all = lv.sum()
    n = {"stores": stores, "calendar": calendar_rows(weeks), "months": month_heads(), "week_heads": week_heads(),
         "wt": wt_all,
         "live": {"first": pd.Timestamp(live.date.min()), "last": pd.Timestamp(live.date.max()),
                  "sales": lv_all.actual / lv_all.runrate - 1, "hour_err": lv_all.hour_err / lv_all.actual,
                  "checks": lv_all.checks / lv_all.checks_expected - 1,
                  "stores": [{"name": i, "sales": r.actual / r.runrate - 1} for i, r in lv.iterrows()]},
         "weeks": weeks, "temps": temps, "rr": rr,
         "n_by_k": {k: [st["name"] for st in stores if st["n"] == k] for k in (2, 3, 4)},
         "fit_lo": min(st["fit_week"] for st in stores), "fit_hi": max(st["fit_week"] for st in stores),
         "n_hol": sum(len(st["hols"]) for st in stores),
         "hol_stores": [st["name"] for st in stores if st["hols"]],
         "no_hol_stores": [st["name"] for st in stores if not st["hols"]],
         "bt": {"weekday": (err[("1 pattern", "weekday")], err[("2 recency", f"weekday, half-life {hl}")]),
                "half": (err[("1 pattern", "weekday x half the type shift by hour")],
                         err[("2 recency", f"weekday x half the type shift by hour, half-life {hl}")]),
                "full": (err[("1 pattern", "weekday x type shift by hour")],
                         err[("2 recency", f"weekday x type shift by hour, half-life {hl}")]),
                "type": err[("1 pattern", "type x weekday")],
                "sp_own": err[("1 special days", "special day: its own pattern")],
                "sp_type": err[("1 special days", "special day: the weekday's pattern with its type's shift")],
                "hl_none": err[("2 recency", "weekday x half the type shift by hour, half-life none")],
                "chk": err[("2 average check", f"check: last {CHECK_MONTHS} months")]},
         "season_shift": SEASON_SHIFT,
         "check": T.check_totals(weeks, temps, rr),
         "last_day": LAST_DAY, "hours_from": HOURS_FROM, "half_life": RECENCY_HALF_LIFE,
         "check_months": CHECK_MONTHS, "min_run": MIN_RUN, "step": STEP,
         "w1": week_start(1), "w52": week_start(N_WEEKS) + pd.Timedelta(days=6), "COLOR": COLOR, "INK_ON": INK_ON}
    n["total_rr"] = float(rr.sum().sum())
    # where summer weeks put more of the day into the afternoon: the busiest type against the quietest
    later = []
    for st in stores:
        reg = [ty for ty in st["types"] if ty["name"] != REMODELED]
        lo, hi = reg[0], reg[-1]
        if hi["parts"][2] - lo["parts"][2] >= 0.005:
            later.append({"name": st["name"], "hi": hi["parts"][2], "lo": lo["parts"][2], "hi_t": hi["name"],
                          "lo_t": lo["name"]})
    n["later"] = sorted(later, key=lambda x: x["lo"] - x["hi"])
    n["live"]["lo"] = min(n["live"]["stores"], key=lambda x: x["sales"])
    n["live"]["hi"] = max(n["live"]["stores"], key=lambda x: x["sales"])
    cor = next(st for st in stores if st["key"] == "Coronado")
    n["coronado_high"] = next(ty for ty in cor["types"] if ty["name"] == "High")
    n["coronado_low"] = next(ty for ty in cor["types"] if ty["name"] == "Low")
    return n


# ------------------------------------------------------------------------------------------
# Workbook
# ------------------------------------------------------------------------------------------
def export_xlsx(n: dict) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold, head_font = Font(bold=True), Font(bold=True, color="FFFFFF")
    ink_fill = PatternFill("solid", fgColor="000000")
    thin = Side(style="thin", color="000000")
    money_fmt, int_fmt = '"$"#,##0', "#,##0"

    def xl(hex_color: str) -> str:
        h = hex_color.lstrip("#").upper()
        return "".join(c * 2 for c in h) if len(h) == 3 else h

    def fill(hex_color: str) -> PatternFill:
        return PatternFill("solid", fgColor=xl(hex_color))

    # read me
    ws = wb.active
    ws.title = "Read me"
    notes = [
        "Parakeet Café · week types for the next 52 weeks: sales and checks by hour for payroll",
        "",
        f"Weeks: the next 52 Monday-to-Sunday weeks (week 1 starts Mon {n['w1']:%b %-d, %Y}; week 52 ends Sun "
        f"{n['w52']:%b %-d, %Y}).",
        "Sales: today's run-rate by day (the run-rate budget: last 45 days of Toast sales to Oct 14, 2026, seasonally "
        "adjusted and held flat; the rest of 2026, then 2027), whole dollars.",
        f"Types: each store's weeks in 2 to 4 types, in runs of at least {MIN_RUN} weeks; each type sells at least "
        f"{STEP:.0%} more than the one below (two types at least). La Jolla Cove: one more type, Remodeled, after its "
        "October 1, 2027 reopening.",
        f"A type's template is the average of its weeks by day of the week; holiday weeks that sell {STEP:.0%} more or "
        "less than their type have a template of their own.",
        f"Hours: each day split along the store's pattern by hour and day of the week in Toast (hour the check was "
        f"opened, Oct 2024 - Oct 14, 2026, recent days weigh more: half-life {RECENCY_HALF_LIFE} days); holidays "
        "along their own pattern from earlier years.",
        f"Checks: each hour's sales divided by the store's average check for that hour and day of the week, last "
        f"{CHECK_MONTHS} months.",
        "",
        "Sheets: Calendar (the type of every week, by store) · Weeks (every week's own run-rate by day against its "
        "template) · Types (each type's weeks and average week) · one sheet per store with its templates (sales, "
        "then "
        "checks; totals are formulas) · All templates (one row per store, template, day and hour) · Check (the "
        "templates applied to the calendar against the run-rate).",
        "To schedule a week: find its type in Calendar, then use that template in the store's sheet. Weeks gives the "
        "week's own sales by day: where a day is a few percent above or below the template, scale that day.",
        f"Live check, {n['live']['first']:%b %-d} - {n['live']['last']:%b %-d, %Y} (after the data): the ten stores "
        f"sold {n['live']['sales']:+.1%} against the run-rate; hourly error {n['live']['hour_err']:.1%}; checks "
        f"{n['live']['checks']:+.1%}.",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(i, 1, t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 140

    # calendar
    ws = wb.create_sheet("Calendar")
    hdr = ["Week", "Monday", "Sunday", "ISO week"] + [st["name"] for st in n["stores"]]
    for j, v in enumerate(hdr, 1):
        c = ws.cell(1, j, v)
        c.font, c.fill, c.alignment = head_font, ink_fill, Alignment(horizontal="center", wrap_text=True)
    for w in range(1, N_WEEKS + 1):
        r = w + 1
        ws.cell(r, 1, w)
        ws.cell(r, 2, week_start(w).date()).number_format = "ddd mmm d, yyyy"
        ws.cell(r, 3, (week_start(w) + pd.Timedelta(days=6)).date()).number_format = "ddd mmm d, yyyy"
        ws.cell(r, 4, int(week_start(w).isocalendar().week))
        for j, st in enumerate(n["stores"], 5):
            sw = st["weeks"]
            t = sw.label[w]
            c = ws.cell(r, j, sw.holiday.get(w, t))
            c.fill = fill(COLOR[t])
            c.font = Font(color=xl(INK_ON[t]), bold=w in sw.holiday)
            c.alignment = Alignment(horizontal="center")
            if w in sw.holiday:
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "E2"
    for j in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 8 if j in (1, 4) else 16
    ws.row_dimensions[1].height = 30

    # every week's own run-rate by day against its template
    ws = wb.create_sheet("Weeks")
    hdr = ["Store", "Week", "Monday", "Template"] + DAYS + ["Week sales", "Template week", "Vs template",
                                                            "Day furthest from the template"]
    for j, v in enumerate(hdr, 1):
        c = ws.cell(1, j, v)
        c.font, c.fill, c.alignment = head_font, ink_fill, Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(n["wt"].itertuples(index=False), 2):
        ws.cell(i, 1, row.store)
        ws.cell(i, 2, row.week)
        ws.cell(i, 3, row.monday).number_format = "ddd mmm d, yyyy"
        c = ws.cell(i, 4, row.template)
        if row.kind == "holiday":
            c.fill, c.font = ink_fill, head_font
        else:
            c.fill, c.font = fill(COLOR[row.template]), Font(color=xl(INK_ON[row.template]))
        for j, d in enumerate(DAYS, 5):
            ws.cell(i, j, round(getattr(row, d))).number_format = money_fmt
        ws.cell(i, 12, f"=SUM(E{i}:K{i})").number_format = money_fmt
        ws.cell(i, 13, int(row.template_sales)).number_format = money_fmt
        ws.cell(i, 14, f"=L{i}/M{i}-1").number_format = '+0.0%;-0.0%;0.0%'
        ws.cell(i, 15, f"{row.largest_day} {row.largest_day_gap:+.0%}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:O{len(n['wt']) + 1}"
    for j, wdt in enumerate([18, 7, 16, 22] + [10] * 7 + [12, 13, 11, 16], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    ws.row_dimensions[1].height = 30

    # types
    ws = wb.create_sheet("Types")
    hdr = ["Store", "Type", "Weeks", "Dates", "Number of weeks", "Average week (sales)", "Checks a week",
           "Holiday weeks with their own template"]
    for j, v in enumerate(hdr, 1):
        c = ws.cell(1, j, v)
        c.font, c.fill = head_font, ink_fill
    r = 2
    for st in n["stores"]:
        for ty in st["types"]:
            ws.cell(r, 1, st["name"])
            c = ws.cell(r, 2, ty["name"])
            c.fill = fill(ty["color"])
            ws.cell(r, 3, "; ".join(x["weeks"] for x in ty["ranges"]))
            ws.cell(r, 4, "; ".join(x["dates"] for x in ty["ranges"]))
            ws.cell(r, 5, ty["n_weeks"])
            ws.cell(r, 6, ty["avg"]).number_format = money_fmt
            ws.cell(r, 7, ty["checks"]).number_format = int_fmt
            ws.cell(r, 8, "; ".join(f"{o['name']} (week {o['week']})" for o in ty["own"]))
            r += 1
    for j, wdt in enumerate([18, 12, 26, 44, 10, 14, 12, 60], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt

    # one sheet per store
    for st in n["stores"]:
        ws = wb.create_sheet(st["name"][:31])
        r = 1
        ws.cell(r, 1, f"{st['name']} · week templates, next 52 weeks (sales in $, then checks; Monday to Sunday)"
                ).font = Font(bold=True, size=13)
        r += 2
        blocks = [(ty["t"], f"{ty['name']} · " + "; ".join(f"{x['weeks']}: {x['dates']}" for x in ty["ranges"]))
                  for ty in st["types"]]
        blocks += [(h["t"], f"{h['name']} · week {h['week']} ({h['dates']}) · instead of the {h['type']} template")
                   for h in st["hols"]]
        for t, title in blocks:
            c = ws.cell(r, 1, title)
            if t.kind == "holiday":
                c.fill, c.font = ink_fill, head_font
            else:
                c.fill, c.font = fill(COLOR[t.name]), Font(bold=True, color=xl(INK_ON[t.name]))
            for what, mat, fmt in (("Sales", t.sales, money_fmt), ("Checks", t.checks, int_fmt)):
                r += 1
                cols = [what] + t.columns + ["Week"]
                for j, v in enumerate(cols, 1):
                    c = ws.cell(r, j, v)
                    c.font, c.border = bold, Border(bottom=thin)
                    c.alignment = Alignment(horizontal="center" if j > 1 else "left")
                top = r + 1
                for i, h in enumerate(t.hours):
                    r += 1
                    ws.cell(r, 1, hour_label(h))
                    for j in range(7):
                        ws.cell(r, j + 2, int(mat[i, j])).number_format = fmt
                    ws.cell(r, 9, f"=SUM(B{r}:H{r})").number_format = fmt
                    ws.cell(r, 9).font = bold
                r += 1
                ws.cell(r, 1, "Day").font = bold
                for j in range(2, 10):
                    col = get_column_letter(j)
                    c = ws.cell(r, j, f"=SUM({col}{top}:{col}{r - 1})")
                    c.number_format, c.font, c.border = fmt, bold, Border(top=thin)
                r += 1
            r += 2
        ws.column_dimensions["A"].width = 10
        for j in range(2, 10):
            ws.column_dimensions[get_column_letter(j)].width = 13

    # long table and check
    ws = wb.create_sheet("All templates")
    lt = T.long_table(n["temps"])
    for j, v in enumerate(lt.columns, 1):
        c = ws.cell(1, j, v)
        c.font, c.fill = head_font, ink_fill
    for i, row in enumerate(lt.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
    ws.freeze_panes = "A2"
    ws = wb.create_sheet("Check")
    ws.cell(1, 1, "Each week's template applied to the calendar, against the run-rate of the 52 weeks "
                  "(difference: rounding to whole dollars per day)").font = bold
    for j, v in enumerate(["Store", "Templates on the calendar", "Run-rate, 52 weeks", "Difference"], 1):
        c = ws.cell(3, j, v)
        c.font, c.fill = head_font, ink_fill
    for i, row in enumerate(n["check"].itertuples(index=False), 4):
        ws.cell(i, 1, row.store)
        ws.cell(i, 2, int(row.templates)).number_format = money_fmt
        ws.cell(i, 3, float(row.runrate)).number_format = '"$"#,##0.00'
        ws.cell(i, 4, f"=B{i}-C{i}").number_format = '"$"#,##0.00'
    for j, wdt in enumerate([20, 26, 22, 14], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    return recalc(XLSX)


def recalc(path: Path, timeout: int = 120) -> dict:
    """Recalculate with LibreOffice (the xlsx skill's script); fail on any formula error."""
    script = next(RECALC.rglob("xlsx/scripts/recalc.py"), None)
    if script is None:
        print("note: recalc.py not found; workbook saved without cached values")
        return {"status": "not recalculated"}
    res = subprocess.run([sys.executable, str(script), str(path), str(timeout)], capture_output=True, text=True)
    try:
        out = json.loads(res.stdout[res.stdout.index("{"):])
    except ValueError:
        raise RuntimeError(f"recalc failed: {res.stdout}\n{res.stderr}")
    if out.get("status") != "success":
        raise RuntimeError(f"workbook has formula errors: {out}")
    return out


# ------------------------------------------------------------------------------------------
# Output
# ------------------------------------------------------------------------------------------
def render(n: dict) -> Path:
    env = Environment(loader=FileSystemLoader(ROOT / "report"), autoescape=False)
    env.filters.update(money=money, fk=fk, fm=fm, pct=pct, pct1=pct1, and_join=and_join)
    html = env.get_template("week_types.html.j2").render(n=n, DAYS=DAYS)
    path = BUILD / "parakeet_week_types_next_52_weeks.html"
    path.write_text(html, encoding="utf-8")
    return path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ds", help="path to the parakeet-cafe-design skill folder")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--shots", help="also write page PNGs to this folder (visual QA)")
    args = ap.parse_args()
    ds = find_design_system(args.ds)
    copy_assets(ds)
    for f in ["photos/hero-1.jpg"]:
        shutil.copy(ds / "assets" / f, BUILD / "assets" / Path(f).name)
    n = collect()
    print(n["check"].to_string(index=False))
    print("workbook:", export_xlsx(n))
    html = render(n)
    print(f"HTML -> {html.relative_to(ROOT)}")
    if not args.no_pdf:
        node_path = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
        env = {**os.environ, "NODE_PATH": node_path}
        cmd = ["node", str(ROOT / "report" / "render_pdf.cjs"), str(html), str(PDF), args.shots or "",
               "13.333in", "7.5in"]
        subprocess.run(cmd, check=True, env=env)
        print(f"PDF  -> {PDF.relative_to(ROOT)}")
        print(f"XLSX -> {XLSX.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
