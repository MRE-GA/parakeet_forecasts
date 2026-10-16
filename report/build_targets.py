"""Build the run-rate review and 2027 manager targets deck, and the targets workbook.

    python report/build_targets.py [--ds PATH_TO_PARAKEET_DESIGN_SYSTEM] [--no-pdf] [--shots DIR]

Part 1 reads the last 45 days of Toast sales (Aug 31 - Oct 14, 2026), seasonally adjusted, as each café's
run-rate (src/runrate_scenario.py): the 2026 finish and 2027 at that level with no further growth,
against the investor budget, and the adjustment that sets each café to its own run-rate. Part 2 sets the
2027 manager targets by month and café: each month of 2026 plus 65% of the growth in our plan (the blend of
our three forecasting models). La Jolla Cove runs that way through September 2027, then on the budget's
remodeled rate from the October 1, 2027 reopening. Nautilus is left out throughout.

Run src/forecast_2027.py, src/hold_scenario.py and src/runrate_scenario.py first. Outputs:
outputs/Parakeet_2027_Run_Rate_and_Targets.pdf and outputs/Parakeet_2027_Manager_Targets.xlsx.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from jinja2 import Environment, FileSystemLoader

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "report"))

import build_deck as bd  # noqa: E402  (evidence per café, formatting, brand assets)
import charts  # noqa: E402
import deck_charts as dc  # noqa: E402
from build_deck import apct, fk, fm, frange, pct, pct0, sfk, sfm, words  # noqa: E402
from build_report import BUILD, copy_assets, find_design_system  # noqa: E402
from forecasting import OUTPUTS, STORE_START, load_daily  # noqa: E402

PDF = OUTPUTS / "Parakeet_2027_Run_Rate_and_Targets.pdf"
XLSX = OUTPUTS / "Parakeet_2027_Manager_Targets.xlsx"
COVE = "La Jolla Cove"
GROWTH_SHARE = 0.65  # targets keep 65% of the growth in our plan (a plan of +10% becomes +6.5%, +20% becomes +13%)
WINDOW = pd.date_range("2026-08-31", "2026-10-14")  # the last 45 days of Toast data
LAG = pd.Timedelta(days=364)  # same weekdays a year earlier
M26 = pd.period_range("2026-01", "2026-12", freq="M")
M27 = pd.period_range("2027-01", "2027-12", freq="M")
FC = pd.period_range("2026-10", "2027-12", freq="M")  # forecast months shown on the charts
BADGE = {"raise": "Raise to run-rate", "lower": "Lower to run-rate", "remodel": "Run-rate + remodel"}
RECALC = Path.home() / ".claude" / "skills"  # the xlsx skill's LibreOffice recalculation script lives here
DPM = 365.25 / 12  # days in an average month: sales per normal day, shown as a monthly figure


# ------------------------------------------------------------------------------------------
# Numbers
# ------------------------------------------------------------------------------------------
def fa(x: float) -> str:
    """Money in millions from $1M up, in thousands below (La Jolla Juicery, La Jolla Cove)."""
    return fm(x) if abs(x) >= 1e6 else fk(x)


def window_45(raw: pd.DataFrame, store: str) -> dict:
    """Sales and checks over the last 45 days and on the same weekdays a year earlier, on days both years
    traded (sums, so cafés can be added up before taking the growth)."""
    d = raw[raw.store == store].set_index("date")
    cur = d.reindex(WINDOW).fillna(0)
    prev = d.reindex(WINDOW - LAG).fillna(0)
    both = (cur.net_sales.to_numpy() > 0) & (prev.net_sales.to_numpy() > 0)
    return {"sales": cur.net_sales.sum(), "s_cur": cur.net_sales.to_numpy()[both].sum(),
            "s_prev": prev.net_sales.to_numpy()[both].sum(), "c_cur": cur.checks.to_numpy()[both].sum(),
            "c_prev": prev.checks.to_numpy()[both].sum()}


def collect() -> dict:
    n0 = bd.collect()  # the budget review's evidence per café (we drop Nautilus)
    raw = load_daily()
    raw_m = raw.assign(month=raw.date.dt.to_period("M")).groupby(["store", "month"]).net_sales.sum()
    rr = pd.read_csv(OUTPUTS / "runrate_45d_2027.csv").set_index("store")
    rrm = pd.read_csv(OUTPUTS / "runrate_45d_2027_by_month.csv", index_col=0)
    rrm.columns = pd.PeriodIndex(rrm.columns, freq="M")
    plan = pd.read_csv(OUTPUTS / "forecast_by_store_month.csv")
    plan = plan.assign(month=pd.PeriodIndex(plan.month, freq="M")).set_index(["store", "month"])[bd.PLAN]
    oct_to_date = raw[raw.date.between("2026-10-01", WINDOW[-1])].groupby("store").net_sales.sum()
    bm = pd.read_csv(ROOT / "data" / "cfo_budget_2027_monthly.csv")
    bm26 = bm[bm.month.str.startswith("2026")].set_index(["store", "month"]).value
    # the run-rate's seasonal pattern summed by month (and actual sales on the same normal days): the base for
    # 'sales without the season'
    sea = pd.read_csv(OUTPUTS / "runrate_45d_seasonal_by_month.csv")
    sea = sea.assign(month=pd.PeriodIndex(sea.month, freq="M")).set_index(["store", "month"])

    nine = n0["nine"]
    stores = nine + [COVE]
    cafes = []
    for s in stores:
        c0 = n0["C"][s]
        c = {k: c0[k] for k in ("key", "name", "short", "ytd26", "ytd_g", "chk_g", "tkt_g", "q1", "q2", "q3",
                                "oct_g", "oct_chk_g", "fy25", "g25", "fy24", "avg_m", "min_m", "max_m", "key_assumption",
                                "b26", "b27", "bg", "q4_b", "plan27", "quarters", "slowest", "slowest_q", "first_q",
                                "note")}
        w45 = window_45(raw, s)
        g45, chk45 = w45["s_cur"] / w45["s_prev"] - 1, w45["c_cur"] / w45["c_prev"] - 1
        r = rr.loc[s]
        # run-rate by whole month, Oct 2026 - Dec 2027: the model alone, so 2026 and 2027 share one seasonality
        run = np.array(rrm.loc[s].reindex(FC), dtype=float)
        # the 2026 estimate by month: October swaps the run-rate for Oct 1-14 for the actual sales
        est26 = run[:3].copy()
        est26[0] += oct_to_date[s] - r.runrate_oct_to_date
        # the budget by month (workbook): Oct-Dec 2026 estimate, then the 2027 objective
        budget = np.asarray(c0["b_months"], dtype=float)
        months26 = np.r_[raw_m.loc[s].reindex(M26[:9]).to_numpy(), est26]
        # 2027 targets, month by month: the same month of 2026 plus a share of the growth in our plan; the Cove
        # switches to the budget's remodeled rate from its reopening
        plan_m = plan.loc[s].reindex(M27).to_numpy()
        target = months26 + GROWTH_SHARE * (plan_m - months26)
        if s == COVE:
            target[9:] = budget[12:15]
        c.update(
            g45=g45, chk45=chk45, w45=w45, to_date=r.actual_2026_to_oct14, rest26=r.runrate_rest_2026,
            rr26=r.runrate_2026E, rr27=r.runrate_2027, rr27_check=r.check_runrate_2027, run=run, budget=budget,
            target=target, months26=months26, target27=target.sum(), plan_m=plan_m, act=raw_m.loc[s],
            # the budget's own 2026 by month (Jan-Sep sales report, Oct-Dec estimate): the base of its 2027 growth
            b26_months=np.array([bm26[(s, f"2026-{i:02d}")] for i in range(1, 13)], dtype=float),
            first_full=pd.Timestamp(STORE_START[s]).to_period("M") + (0 if pd.Timestamp(STORE_START[s]).day == 1 else 1),
        )
        assert abs(c["b26_months"].sum() - c["b26"]) < 1, s
        assert abs(est26.sum() + c["ytd26"] - c["rr26"]) < 1, s  # Jan-Sep + Oct-Dec = run-rate 2026 finish
        assert abs(months26.sum() - c["rr26"]) < 1, s
        assert abs(run[3:].sum() - c["rr27"]) < 1, s
        # the proposed 2027 objective, by month: the run-rate; the Cove switches to the budget's remodeled rate
        # from its October 1, 2027 reopening
        path = run.copy()
        if s == COVE:
            path[12:15] = budget[12:15]
            c["action"] = "remodel"
        else:
            c["action"] = "raise" if c["rr27"] > c["b27"] else "lower"
        c["path"] = path
        c["prop27"] = path[3:].sum()
        c["d26"] = c["rr26"] / c["b26"] - 1
        c["adj"] = c["prop27"] - c["b27"]
        c["need"] = c["b27"] / c["prop27"] - 1  # growth over today's level the budget needs
        c["rr_g"] = c["rr27"] / c["rr26"] - 1  # 2027 at today's level against the 2026 finish
        c["prop_g"] = c["prop27"] / c["rr26"] - 1
        c["tg"] = c["target27"] / c["rr26"] - 1
        c["t_vs_rr"] = c["target27"] - c["prop27"]
        c["t_vs_b"] = c["target27"] - c["b27"]
        c["ask"] = c["t_vs_rr"] / c["prop27"]  # what the target asks on top of today's level
        c["quarters_t"] = [target[i:i + 3].sum() for i in (0, 3, 6, 9)]
        qf = fk if max(c["quarters_t"]) < 1e6 else fm
        c["quarters_s"] = [qf(q) for q in c["quarters_t"]]
        # sales without the season, as a monthly equivalent: history (actual sales on normal days over the
        # seasonal factors of those days), the factor sums for the run-rate months, and today's level
        se = sea.loc[s]
        f_hist = se.factor_sum.loc[:pd.Period("2026-09", "M")]
        c["sa_hist"] = se.sales.loc[f_hist.index] / f_hist.where(f_hist > 0) * DPM
        c["ff"] = se.factor_sum.reindex(FC).to_numpy()
        c["level_m"] = c["rr27"] / c["ff"][3:].sum() * DPM
        assert np.allclose(run, c["level_m"] / DPM * c["ff"], rtol=1e-4), s  # run-rate = today's level x seasons
        cafes.append(c)
    C = {c["key"]: c for c in cafes}
    n = {"cafes": cafes, "C": C, "nine": nine, "stores": stores,
         "first_day": n0["first_day"], "last_day": n0["last_day"], "store_days": n0["store_days"]}

    # ---- totals: the nine existing cafés, and with La Jolla Cove ------------------------------
    def total(group: list[str]) -> dict:
        g = [C[s] for s in group]
        out = {k: sum(c[k] for c in g) for k in ("ytd26", "to_date", "rest26", "rr26", "rr27", "rr27_check", "b26",
                                                   "b27", "q4_b", "prop27", "target27", "plan27", "fy25")}
        for k in ("run", "path", "budget", "target", "months26", "plan_m"):
            out[k] = sum(c[k] for c in g)
        out["act"] = sum(c["act"] for c in g)
        w = {k: sum(c["w45"][k] for c in g) for k in g[0]["w45"]}
        out["sales45"] = w["sales"]
        out["g45"] = w["s_cur"] / w["s_prev"] - 1
        out["chk45"] = w["c_cur"] / w["c_prev"] - 1
        out["d26"] = out["rr26"] / out["b26"] - 1
        out["adj"] = out["prop27"] - out["b27"]
        out["rr_g"] = out["rr27"] / out["rr26"] - 1
        out["prop_g"] = out["prop27"] / out["rr26"] - 1
        out["bg"] = out["b27"] / out["b26"] - 1
        out["tg"] = out["target27"] / out["rr26"] - 1
        out["t_vs_rr"] = out["target27"] - out["prop27"]
        out["t_vs_b"] = out["target27"] - out["b27"]
        out["ask"] = out["t_vs_rr"] / out["prop27"]
        out["booked"] = out["to_date"]
        out["booked_share"] = out["to_date"] / out["rr26"]
        out["quarters_t"] = [out["target"][i:i + 3].sum() for i in (0, 3, 6, 9)]
        out["quarters_26"] = [out["months26"][i:i + 3].sum() for i in (0, 3, 6, 9)]
        out["quarters_g"] = [a / b - 1 for a, b in zip(out["quarters_t"], out["quarters_26"])]
        out["plan_g"] = out["plan_m"].sum() / out["rr26"] - 1  # growth in our plan, before the targets' share
        return out

    n["t9"] = total(nine)
    n["t10"] = total(stores)
    up = [C[s] for s in nine if C[s]["adj"] > 0]
    down = [C[s] for s in nine if C[s]["adj"] < 0]
    n["up"] = sorted(up, key=lambda c: -c["adj"])
    n["down"] = sorted(down, key=lambda c: c["adj"])
    n["up_total"] = sum(c["adj"] for c in up)
    n["down_total"] = sum(c["adj"] for c in down)
    by_ask = sorted((C[s] for s in nine), key=lambda c: c["ask"])
    n["ask_lo"], n["ask_hi"] = by_ask[0], by_ask[-2:][::-1]
    n["qtr"] = quarterly_same_store([C[s] for s in nine])
    for c in cafes:
        c.update(story(c))
    check(n)
    return n


def quarterly_same_store(group: list[dict]) -> dict:
    """Same-store growth against the same quarter (and year) a year earlier, 2024-2027, for the slide on growth by
    quarter. Monthly sales: actual to September 2026, the Oct-Dec 2026 estimate (actual to Oct 14, then the
    run-rate), then 2027 at today's run-rate, or the target. A café counts in a period once it traded the whole
    period a year earlier; the budget's 2027 is compared with the budget's own 2026."""
    def series(c, kind):
        tail = np.r_[c["months26"][9:12], c["run"][3:] if kind == "run" else c["target"]]
        return pd.concat([c["act"].loc[:pd.Period("2026-09", "M")],
                          pd.Series(tail, index=pd.period_range("2026-10", "2027-12", freq="M"))])
    ser = {kind: {c["key"]: series(c, kind) for c in group} for kind in ("run", "target")}

    def growth(months, kind="run"):
        ok = [c for c in group if months[0] - 12 >= c["first_full"]]
        cur = sum(ser[kind][c["key"]].reindex(months).sum() for c in ok)
        prev = sum(ser[kind][c["key"]].reindex(months - 12).sum() for c in ok)
        return cur / prev - 1, len(ok)

    qmonths = lambda y, q: pd.period_range(f"{y}-{3 * q - 2:02d}", periods=3, freq="M")  # noqa: E731
    out = {"quarters": {y: [growth(qmonths(y, q)) for q in (1, 2, 3, 4)] for y in (2024, 2025, 2026, 2027)},
           "years": {y: growth(pd.period_range(f"{y}-01", periods=12, freq="M")) for y in (2024, 2025, 2026, 2027)},
           "target": [growth(qmonths(2027, q), "target")[0] for q in (1, 2, 3, 4)],
           "budget": [sum(c["budget"][3 * q:3 * q + 3].sum() for c in group)
                      / sum(c["b26_months"][3 * q - 3:3 * q].sum() for c in group) - 1 for q in (1, 2, 3, 4)]}
    out["target_year"] = sum(c["target27"] for c in group) / sum(c["rr26"] for c in group) - 1
    out["budget_year"] = sum(c["b27"] for c in group) / sum(c["b26"] for c in group) - 1
    # the quarters since late 2024, when most cafés compare: actual, then the Q4 2026 estimate
    since = [g for y in (2024, 2025, 2026) for q, (g, _) in enumerate(out["quarters"][y], 1) if (y, q) >= (2024, 4)]
    out["lo"], out["hi"] = min(since), max(since)
    return out


def check(n: dict) -> None:
    """Fail loudly if refreshed numbers no longer support a sentence on the pages."""
    C, t9 = n["C"], n["t9"]
    nine = [C[s] for s in n["nine"]]
    assert abs(t9["d26"]) < 0.01 and abs(t9["adj"] / t9["b27"]) < 0.01  # "matches the budget"
    assert sum(c["rr26"] < c["b26"] for c in nine) >= 5 and all(abs(c["d26"]) < 0.025 for c in nine)  # "a little
    # lower for most cafés"
    assert abs(t9["bg"] - t9["rr_g"]) < 0.005  # "both are about X% above 2026"
    assert all(c["t_vs_rr"] > 0 for c in n["cafes"])  # targets sit above today's level ("asks for ... more")
    assert C["Brentwood"]["max_m"] / C["Brentwood"]["min_m"] < 1.15  # "held at about ... since January"
    assert C["Coronado"]["q1"] > C["Coronado"]["q2"] > C["Coronado"]["q3"]  # "cooled every quarter"
    assert C["La Jolla"]["q2"] < C["La Jolla"]["q3"]  # "slowed in the spring, picked up again"
    assert C["Little Italy"]["chk45"] > 0.75 * C["Little Italy"]["g45"]  # "almost all from more guests"
    assert all(abs(C[s]["rr27_check"] / C[s]["rr27"] - 1) < 0.05 for s in n["nine"])  # two methods within 5%
    bh = C["Beverly Hills"]
    ticket = (1 + bh["g45"]) / (1 + bh["chk45"]) - 1
    assert abs(np.log1p(ticket) - np.log1p(bh["chk45"])) < 0.05  # "about half guests, half average check"
    assert C["Carlsbad"]["chk45"] > 0.85 * C["Carlsbad"]["g45"]  # "nearly all from more guests"
    jj = C["La Jolla Juicery"]
    sa = jj["sa_hist"]
    assert sa.loc["2026-01":"2026-09"].min() > sa.loc["2025-01":"2025-12"].max()  # "stepped up in January 2026"
    assert abs(jj["level_m"] / sa.loc["2026-01":"2026-09"].mean() - 1) < 0.03  # "today's level is its 2026 average"
    d = jj["run"][3:] - jj["months26"]  # 2027 at today's level minus 2026, by month
    assert d[2] + d[7] < jj["rr27"] - jj["rr26"] < 0 and np.delete(d, [2, 7]).sum() > -2000  # "repeats 2026 except
    # March and August"
    assert abs(n["up_total"] + n["down_total"] - t9["adj"]) < 1
    assert abs(t9["tg"] - GROWTH_SHARE * t9["plan_g"]) < 1e-9  # the nine's targets keep 65% of the plan's growth
    q = n["qtr"]
    run27 = [g for g, _ in q["quarters"][2027]]
    assert all(a > b for a, b in zip(run27, run27[1:])) and abs(run27[-1]) < 0.015  # "fades to about zero"
    assert abs(q["years"][2026][0] - (t9["rr26"] / t9["fy25"] - 1)) < 1e-6  # 2026E: all nine cafés
    assert abs(q["years"][2027][0] - t9["rr_g"]) < 1e-9 and abs(q["target_year"] - t9["tg"]) < 1e-9
    assert max(q["budget"]) - min(q["budget"]) < 0.005  # "the same every quarter"


def story(c: dict) -> dict:
    """Headlines and plain-English text for the café's run-rate page (part 1) and target page (part 2)."""
    s, name = c["key"], c["name"]
    o = {}
    if c["action"] == "raise":
        o["h1"] = f"{name}: today's level already beats the 2027 budget by {fk(c['adj'])}"
    elif c["action"] == "lower":
        o["h1"] = f"{name}: the 2027 budget needs {apct(c['need'])} growth over today's level"
    else:
        o["h1"] = ("La Jolla Cove: with the remodel, today's level matches the budget" if abs(c["adj"]) < 10_000 else
                   f"La Jolla Cove: with the remodel, today's level lands {fk(c['adj'])} "
                   f"{'under' if c['adj'] < 0 else 'over'} the budget")
    level = (f"Held at its last-45-day level with normal seasons, {name} sells {fa(c['rr27'])} in 2027, "
             f"{fk(c['adj'])} {'more' if c['adj'] > 0 else 'less'} than the budget's {fa(c['b27'])}")
    trend = f"sales are {pct(c['g45'])} on the same 45 days last year, with checks {pct(c['chk45'])}"
    if s == "Beverly Hills":
        o["found"] = (f"{level}. It is still accelerating: {pct(c['ytd_g'])} for January–September and "
                      f"{pct(c['g45'])} in the last 45 days, with checks {pct(c['chk45'])}. About half the growth is "
                      f"more guests, half a higher average check.")
    elif s == "Brentwood":
        o["found"] = (f"Brentwood has held at about {fk(c['avg_m'])} a month since January (between {fk(c['min_m'])} "
                      f"and {fk(c['max_m'])}). {level}; the budget's {pct(c['bg'])} already allows for Leora opening "
                      f"next door. Its {pct(c['g45'])} on last year is still last year's ramp being lapped.")
    elif s == "Carlsbad":
        o["found"] = f"{level}. Carlsbad keeps growing at a steady pace: {trend}, nearly all from more guests."
    elif s == "Coronado":
        o["found"] = (f"Coronado's growth has cooled every quarter of 2026 ({pct0(c['q1'])}, {pct0(c['q2'])}, "
                      f"{pct0(c['q3'])}); {trend}. {level}: the budget needs {apct(c['need'])} growth over "
                      f"today's level.")
    elif s == "Carmel Valley":
        o["found"] = (f"{level}. {trend[0].upper() + trend[1:]}. The budget's {pct(c['bg'])} books a hit from "
                      f"Goop Kitchen next door that the sales do not show so far.")
    elif s == "La Jolla":
        o["found"] = (f"La Jolla Café slowed in the spring ({pct0(c['q2'])} in Q2) and picked up again "
                      f"({pct0(c['q3'])} in Q3); {trend}. {level}: the budget needs {apct(c['need'])} growth over "
                      f"today's level.")
    elif s == "La Jolla Juicery":
        yo = lambda m: c["act"][pd.Period(f"2026-{m:02d}", "M")] / c["act"][pd.Period(f"2025-{m:02d}", "M")] - 1  # noqa: E731
        o["found"] = (f"The Juicery stepped up in January 2026 and has held that level since, so every month beats "
                      f"2025 ({pct(c['ytd_g'])} for January–September) but today's level is its 2026 average. Held "
                      f"there, 2027 repeats 2026 month by month except March and August, two unusually strong months "
                      f"({pct0(yo(3))} and {pct0(yo(8))} on last year): {fa(c['rr27'])}, {pct(c['rr_g'])} on 2026 and "
                      f"{fk(c['adj'])} less than the budget's {fa(c['b27'])}, which needs {apct(c['need'])} growth "
                      f"over today's level.")
    elif s == "Little Italy":
        o["found"] = (f"Little Italy stepped up in March and is still climbing: {trend}, so almost all the growth "
                      f"is more guests. {level}.")
    elif s == "Newport":
        o["found"] = (f"Newport has settled into steady growth; {trend}. {level}: the budget needs "
                      f"{apct(c['need'])} growth over today's level.")
    elif s == COVE:
        js = c["run"][3:12].sum()
        o["found"] = (f"The hotel around the Cove has been under remodel since March 2026, so the café is open but "
                      f"soft ({pct(c['ytd_g'])} for January–September). At today's level it sells {fa(js)} in "
                      f"January–September 2027 ({fa(c['rr27'])} for the full year); adding the budget's $125K a month "
                      f"from the October 1, 2027 reopening gives {fa(c['prop27'])}, "
                      + (f"in line with the budget's {fa(c['b27'])}." if abs(c["adj"]) < 10_000 else
                         f"{fk(c['adj'])} {'under' if c['adj'] < 0 else 'over'} the budget's {fa(c['b27'])}."))
    o["proposal"] = (f"Proposed 2027 objective: {fa(c['prop27'])}, {sfk(c['adj'])} against the budget "
                     + ("(run-rate to September, then the remodeled $125K a month)." if s == COVE else
                        "(today's run-rate, no further growth)."))
    # part 2
    o["t_h1"] = f"{name}: 2027 target {fa(c['target27'])}, {pct(c['tg'])} on 2026"
    # the summary line under the target page
    ask = f"{apct(c['ask'])} over today's run-rate (about {fk(c['t_vs_rr'] / 12)} a month)"
    if s == COVE:
        js_ask = c["target"][:9].sum() / c["run"][3:12].sum() - 1
        o["t_line"] = (f"{fa(c['target27'])} for 2027: 65% of the plan's growth to September ({apct(js_ask)} over "
                       f"today's run-rate), then the remodeled $125K a month from the October 1 reopening; "
                       f"{sfk(c['t_vs_b'])} against the budget.")
    elif s == "Brentwood":
        o["t_line"] = (f"{fa(c['target27'])} for 2027 asks Brentwood to grow again after a flat 2026: {ask}, "
                       f"{sfk(c['t_vs_b'])} against a budget that already allows for Leora.")
    else:
        o["t_line"] = (f"{fa(c['target27'])} for 2027: {pct(c['tg'])} on 2026, {ask} and {sfk(c['t_vs_b'])} against "
                       f"the budget.")
    return o


# ------------------------------------------------------------------------------------------
# Charts
# ------------------------------------------------------------------------------------------
def year_labels(months: list[pd.Period], labels: dict[int, str]) -> list[tuple]:
    out = []
    for y in sorted({m.year for m in months}):
        idx = [i for i, m in enumerate(months) if m.year == y]
        out.append((idx[0] - 0.5, idx[-1] + 0.5, labels[y]))
    return out


def chart_start(store: str) -> pd.Period:
    start = pd.Timestamp(STORE_START[store])
    return max(pd.Period("2024-01", "M"), start.to_period("M") + (0 if start.day == 1 else 1))


def growth_vs_last_year(group: list[dict], months_act, g45: float) -> dict:
    """Growth vs the same month a year earlier, for one café or a group of cafés, for the strip under the
    run-rate charts. Actual months are same-store: a café counts in a month only if it traded the whole month
    a year earlier. Forecast months compare the run-rate (or the budget) with the same month a year earlier:
    actual sales, or the run-rate for Oct-Dec 2026; the budget's 2027 is compared with the budget's own 2026."""
    def actual(p):
        ok = [c for c in group if p - 12 >= c["first_full"]]
        if not ok:
            return np.nan
        return sum(c["act"].get(p, 0.0) for c in ok) / sum(c["act"].get(p - 12, 0.0) for c in ok) - 1

    def base(c, q):  # a café's sales in an earlier month: actual, or the run-rate after September 2026
        return c["run"][FC.get_loc(q)] if q >= FC[0] else c["act"][q]

    run = [sum(c["path"][i] for c in group) / sum(base(c, p - 12) for c in group) - 1 for i, p in enumerate(FC)]
    budget = [sum(c["budget"][i] for c in group)
              / sum(c["act"][p - 12] if p.year == 2026 else c["b26_months"][p.month - 1] for c in group) - 1
              for i, p in enumerate(FC)]
    return {"actual": [actual(p) for p in months_act], "run": run, "budget": budget, "g45": g45}


def trend_parts(group: list[dict], months_act) -> tuple[dict, dict]:
    """The two series behind the Part 2 chart, for one café or a group: monthly sales, and the same without the
    season (each café's sales per normal day as a monthly equivalent, added up across the group). Oct-Dec 2026:
    the estimate on top (actual to Oct 14, then the run-rate), the run-rate alone below."""
    def total(f):
        return np.sum([f(c) for c in group], axis=0)

    def budget_level(c):  # the budget's level, season removed: for the year, or (Cove) before and after the reopening
        out = np.empty(12)
        for a, b in ([(3, 12), (12, 15)] if c["key"] == COVE else [(3, 15)]):
            out[a - 3:b - 3] = c["budget"][a:b].sum() / c["ff"][a:b].sum() * DPM
        return out

    monthly = {"actual": total(lambda c: c["act"].reindex(months_act).to_numpy(dtype=float)),
               "est26": total(lambda c: c["months26"][9:12]), "target": total(lambda c: c["target"]),
               "run": total(lambda c: c["path"][3:]), "budget": total(lambda c: c["budget"][3:])}
    trend = {"actual": total(lambda c: c["sa_hist"].reindex(months_act).to_numpy(dtype=float)),
             "est26": total(lambda c: c["run"][:3] / c["ff"][:3] * DPM),
             "target": total(lambda c: c["target"] / c["ff"][3:] * DPM),
             "run": total(lambda c: c["path"][3:] / c["ff"][3:] * DPM), "budget": total(budget_level)}
    return monthly, trend


def target_year_labels(months, act: pd.Series, first_full: pd.Period, rr26: float, target27: float, tg: float,
                       prop27: float, b27: float, cove: bool = False) -> list[tuple]:
    """Year labels for the Part 2 charts: each year's sales with its growth on the year before (when both years
    were traded in full), then the 2027 target against today's run-rate and the budget."""
    full = lambda y: first_full <= pd.Period(f"{y}-01", "M")  # noqa: E731
    tot = {y: act[act.index.year == y].sum() for y in range(2022, 2026)}
    tot[2026] = rr26
    labels = {}
    for y in sorted({m.year for m in months}):
        if y == 2027:
            labels[y] = (f"2027 target {fa(target27)}, {pct(tg)}\n"
                         f"{'with remodel' if cove else 'run-rate'} {fa(prop27)} · budget {fa(b27)}")
        elif not full(y):
            labels[y] = f"{y}\nfrom {months[0].strftime('%b')}"
        else:
            g = f", {pct(tot[y] / tot[y - 1] - 1)}" if full(y - 1) and tot.get(y - 1, 0) > 0 else ""
            labels[y] = f"{'2026E' if y == 2026 else y}\n{fa(tot[y])}{g}"
    return year_labels(months, labels)


def make_charts(n: dict) -> dict:
    out = {"cafe": {}}
    t9 = n["t9"]
    # the nine cafés: actuals from 2025, run-rate, budget
    act9 = t9["act"].loc[pd.Period("2025-01", "M"):pd.Period("2026-09", "M")]
    months = list(act9.index) + list(FC)
    out["brand_rr"] = dc.runrate_monthly(
        [m.to_timestamp() for m in act9.index], act9.to_numpy(), [m.to_timestamp() for m in FC], t9["run"],
        t9["budget"], year_labels(months, {2025: f"2025\n{fm(t9['fy25'])}", 2026: f"2026E\n{fm(t9['rr26'])}",
                                           2027: f"2027\nrun-rate {fm(t9['rr27'])} · budget {fm(t9['b27'])}"}),
        money="m", growth=growth_vs_last_year([n["C"][s] for s in n["nine"]], act9.index, t9["g45"]),
        size=(7.4, 4.0), name="brand-rr")
    nine = [n["C"][s] for s in n["nine"]]
    q = n["qtr"]["quarters"]
    out["qtrend"] = dc.quarter_trend(
        [2024, 2025, 2026, 2027], [g for y in (2024, 2025) for g, _ in q[y]] + [g for g, _ in q[2026][:3]],
        q[2026][3][0], [g for g, _ in q[2027]], n["qtr"]["budget"], n["qtr"]["target"], name="qtrend")
    adj = sorted(nine, key=lambda c: -c["adj"])
    out["adjust"] = dc.gap_bars([c["short"] for c in adj], [c["adj"] for c in adj],
                                [charts.PINK if c["adj"] > 0 else charts.G400 for c in adj], size=(5.2, 4.2))
    out["targets_dots"] = dc.growth_dots(
        [c["short"] for c in nine],
        [("Budget", [c["bg"] for c in nine], "D", dc.BUDGET),
         ("Today's run-rate", [c["rr_g"] for c in nine], "o", charts.INK),
         ("2027 target", [c["tg"] for c in nine], "o", charts.PINK)],
        size=(7.2, 4.0), name="targets-dots")
    mo, tr = trend_parts(nine, act9.index)
    out["targets9"] = dc.target_trend(
        months, mo, tr, target_year_labels(months, t9["act"], max(c["first_full"] for c in nine), t9["rr26"],
                                           t9["target27"], t9["tg"], t9["prop27"], t9["b27"]),
        money="m", name="targets-nine")
    for c in n["cafes"]:
        act = c["act"].loc[chart_start(c["key"]):pd.Period("2026-09", "M")]
        months = list(act.index) + list(FC)
        cove = c["key"] == COVE
        labels = {2024: (f"2024\n{fa(c['fy24'])}" if pd.Timestamp(STORE_START[c["key"]]) <= pd.Timestamp("2024-01-01")
                         else f"2024\nfrom {act.index[0].strftime('%b')}"),
                  2025: f"2025\n{fa(c['fy25'])}", 2026: f"2026E\n{fa(c['rr26'])}",
                  2027: f"2027\n{'with remodel' if cove else 'run-rate'} {fa(c['prop27'])} · budget {fa(c['b27'])}"}
        rr_chart = dc.runrate_monthly([m.to_timestamp() for m in act.index], act.to_numpy(),
                                      [m.to_timestamp() for m in FC], c["path"], c["budget"],
                                      year_labels(months, labels),
                                      label27="Run-rate 2027 + remodel" if cove else "Run-rate 2027 (no growth)",
                                      growth=growth_vs_last_year([c], act.index, c["g45"]),
                                      size=(7.5, 4.1), name=f"rr-{c['key']}")
        mo, tr = trend_parts([c], act.index)
        tt = dc.target_trend(months, mo, tr,
                             target_year_labels(months, c["act"], c["first_full"], c["rr26"], c["target27"], c["tg"],
                                                c["prop27"], c["b27"], cove=cove),
                             run_label="Run-rate + remodel" if cove else "Today's run-rate", name=f"tt-{c['key']}")
        out["cafe"][c["key"]] = {"rr": rr_chart, "targets": tt}
    return out


# ------------------------------------------------------------------------------------------
# Workbook: monthly targets and tracking
# ------------------------------------------------------------------------------------------
def export_xlsx(n: dict) -> dict:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    arial = "Arial"
    title, bold, normal = Font(name=arial, size=14, bold=True), Font(name=arial, size=10, bold=True), Font(name=arial,
                                                                                                          size=10)
    blue = Font(name=arial, size=10, color="0000FF")
    green = Font(name=arial, size=10, color="008000")
    muted = Font(name=arial, size=9, italic=True, color="565A5C")
    head_fill = PatternFill("solid", fgColor="F2F2F2")
    input_fill = PatternFill("solid", fgColor="FFF8DC")
    thin = Side(style="thin", color="000000")
    money, pct_fmt = '$#,##0;($#,##0);"-"', '0.0%;(0.0%);"-"'
    stores = n["cafes"]
    months = [m.strftime("%b %Y") for m in M27]

    # ---- Targets 2027 ------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Targets 2027"
    ws["A1"], ws["A1"].font = "Parakeet Café · 2027 sales targets by store", title
    ws["A2"], ws["A2"].font = ("Net sales targets per month (each month of 2026 plus 65% of the growth in our plan; "
                               "La Jolla Cove switches to the budget's remodeled "
                               "rate from its Oct 1, 2027 reopening). Blue = values from the analysis; black = formulas."
                               ), muted
    hdr = ["Store"] + months + ["Q1", "Q2", "Q3", "Q4", "Total 2027", "2026 finish (run-rate)", "Growth"]
    for j, h in enumerate(hdr, start=1):
        cell = ws.cell(4, j, h)
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(horizontal="center" if j > 1 else "left", wrap_text=True)
        cell.border = Border(bottom=thin)
    first = 5
    rows = {}
    order = [c for c in stores if c["key"] != COVE] + [None] + [c for c in stores if c["key"] == COVE] + [None]
    r = first
    for c in order:
        if c is None:
            r += 1  # placeholder rows for the subtotal and total, filled below
            continue
        rows[c["key"]] = r
        ws.cell(r, 1, c["name"]).font = normal
        for i, v in enumerate(c["target"]):
            cell = ws.cell(r, 2 + i, round(float(v), 2))
            cell.font, cell.number_format = blue, money
        r += 1
    nine_rows = [rows[c["key"]] for c in stores if c["key"] != COVE]
    sub_row = max(nine_rows) + 1
    cove_row = rows[COVE]
    tot_row = cove_row + 1
    ws.cell(sub_row, 1, "Nine existing cafés").font = bold
    ws.cell(tot_row, 1, "Total, ten stores").font = bold
    col = {name: j for j, name in enumerate(hdr, start=1)}
    q_cols = {"Q1": (2, 4), "Q2": (5, 7), "Q3": (8, 10), "Q4": (11, 13)}
    for rr_ in nine_rows + [cove_row]:
        for q, (a, b) in q_cols.items():
            ws.cell(rr_, col[q], f"=SUM({get_column_letter(a)}{rr_}:{get_column_letter(b)}{rr_})")
        ws.cell(rr_, col["Total 2027"], f"=SUM(B{rr_}:M{rr_})")
        key = next(k for k, v in rows.items() if v == rr_)
        cell = ws.cell(rr_, col["2026 finish (run-rate)"], round(float(n["C"][key]["rr26"]), 2))
        cell.font = blue
        tcol, bcol = get_column_letter(col["Total 2027"]), get_column_letter(col["2026 finish (run-rate)"])
        ws.cell(rr_, col["Growth"], f"=IF({bcol}{rr_}=0,\"\",{tcol}{rr_}/{bcol}{rr_}-1)")
    for j in range(2, col["2026 finish (run-rate)"] + 1):
        L = get_column_letter(j)
        ws.cell(sub_row, j, f"=SUM({L}{min(nine_rows)}:{L}{max(nine_rows)})").font = bold
        ws.cell(tot_row, j, f"={L}{sub_row}+{L}{cove_row}").font = bold
    for rr_ in (sub_row, tot_row):
        tcol, bcol = get_column_letter(col["Total 2027"]), get_column_letter(col["2026 finish (run-rate)"])
        ws.cell(rr_, col["Growth"], f"=IF({bcol}{rr_}=0,\"\",{tcol}{rr_}/{bcol}{rr_}-1)").font = bold
        for j in range(1, len(hdr) + 1):
            ws.cell(rr_, j).border = Border(top=thin)
    for rr_ in range(first, tot_row + 1):
        for j in range(2, len(hdr) + 1):
            ws.cell(rr_, j).number_format = pct_fmt if j == col["Growth"] else money
            if ws.cell(rr_, j).font != bold and ws.cell(rr_, j).font != blue:
                ws.cell(rr_, j).font = normal if rr_ not in (sub_row, tot_row) else bold
    ws.cell(cove_row, 11).comment = Comment("Oct-Dec 2027: the budget's remodeled rate, $125K a month "
                                            "(reopening Oct 1, 2027).", "Analysis")
    ws.column_dimensions["A"].width = 22
    for j in range(2, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 12.5
    ws.freeze_panes = "B5"

    # ---- Monthly tracking ------------------------------------------------------------------------
    tr = wb.create_sheet("Monthly tracking")
    tr["A1"], tr["A1"].font = "Monthly tracking: enter each month's actual net sales in the cream cells", title
    tr["A2"], tr["A2"].font = ("Targets link to 'Targets 2027' (green). Variances fill in once an actual is entered."
                               ), muted
    th = ["Store", "Month", "Target", "Actual", "Variance $", "Variance %"]
    for j, h in enumerate(th, start=1):
        cell = tr.cell(4, j, h)
        cell.font, cell.fill, cell.border = bold, head_fill, Border(bottom=thin)
    r = 5
    for c in stores:
        src = rows[c["key"]]
        for i, m in enumerate(months):
            tr.cell(r, 1, c["name"]).font = normal
            tr.cell(r, 2, m).font = normal
            cell = tr.cell(r, 3, f"='Targets 2027'!{get_column_letter(2 + i)}{src}")
            cell.font, cell.number_format = green, money
            cell = tr.cell(r, 4)
            cell.fill, cell.font, cell.number_format = input_fill, blue, money
            tr.cell(r, 5, f'=IF(D{r}="","",D{r}-C{r})').number_format = money
            tr.cell(r, 6, f'=IF(D{r}="","",D{r}/C{r}-1)').number_format = pct_fmt
            for j in (5, 6):
                tr.cell(r, j).font = normal
            r += 1
    last_tr = r - 1
    # how to fill in, with a worked example outside the tracked range (so it is never counted)
    tr["H4"], tr["H4"].font = "How to fill in", bold
    for i, text in enumerate([
        "1. When a month closes, type the store's actual net sales from Toast, in dollars, in its cream cell "
        "(column D).",
        "2. Variance $ and % fill in on their own; the 'Year to date' tab adds up the months entered.",
        "3. Leave months blank until they close. Targets (green) link to 'Targets 2027': change them there.",
        "Colors: blue = typed values · black = formulas · green = links to another tab.",
    ], start=5):
        tr.cell(i, 8, text).font = normal
    tr["H10"], tr["H10"].font = "Example (not counted)", bold
    for j, h in enumerate(th, start=8):
        cell = tr.cell(11, j, h)
        cell.font, cell.fill, cell.border = bold, head_fill, Border(bottom=thin)
    tr["H12"], tr["I12"] = stores[0]["name"], months[0]
    tr["J12"] = "=C5"  # the first tracked row: the first store's January target
    tr["K12"] = round(float(stores[0]["target"][0]) * 1.02, -3)  # an actual about 2% over target
    tr["L12"], tr["M12"] = "=IF(K12=\"\",\"\",K12-J12)", "=IF(K12=\"\",\"\",K12/J12-1)"
    for L, fnt, fmt in (("H", muted, None), ("I", muted, None), ("J", green, money), ("K", blue, money),
                        ("L", normal, money), ("M", normal, pct_fmt)):
        tr[f"{L}12"].font = fnt
        if fmt:
            tr[f"{L}12"].number_format = fmt
    tr["K12"].fill = input_fill
    tr.column_dimensions["A"].width = 22
    for L in "BCDEF":
        tr.column_dimensions[L].width = 14
    tr.column_dimensions["G"].width = 4
    tr.column_dimensions["H"].width = 16
    for L in "IJKLM":
        tr.column_dimensions[L].width = 13
    tr.freeze_panes = "A5"

    # ---- Summary (year to date) ------------------------------------------------------------------
    sm = wb.create_sheet("Year to date", 1)
    sm["A1"], sm["A1"].font = "Year to date against target", title
    sm["A2"], sm["A2"].font = "Calculated from 'Monthly tracking': only months with an actual entered count.", muted
    sh = ["Store", "Target 2027", "Target to date", "Actual to date", "Variance $", "Variance %", "Months entered"]
    for j, h in enumerate(sh, start=1):
        cell = sm.cell(4, j, h)
        cell.font, cell.fill, cell.border = bold, head_fill, Border(bottom=thin)
    rng = lambda L: f"'Monthly tracking'!${L}$5:${L}${last_tr}"  # noqa: E731
    r = 5
    for c in stores:
        sm.cell(r, 1, c["name"]).font = normal
        sm.cell(r, 2, f"='Targets 2027'!{get_column_letter(col['Total 2027'])}{rows[c['key']]}").font = green
        sm.cell(r, 3, f'=SUMIFS({rng("C")},{rng("A")},A{r},{rng("D")},"<>")')
        sm.cell(r, 4, f"=SUMIFS({rng('D')},{rng('A')},A{r})")
        sm.cell(r, 5, f'=IF(G{r}=0,"",D{r}-C{r})')
        sm.cell(r, 6, f'=IF(G{r}=0,"",D{r}/C{r}-1)')
        sm.cell(r, 7, f'=COUNTIFS({rng("A")},A{r},{rng("D")},"<>")')
        for j in range(3, 8):
            sm.cell(r, j).font = normal
        r += 1
    sm.cell(r, 1, "Total, ten stores").font = bold
    for j, L in ((2, "B"), (3, "C"), (4, "D")):
        sm.cell(r, j, f"=SUM({L}5:{L}{r - 1})").font = bold
    sm.cell(r, 5, f'=IF(D{r}=0,"",D{r}-C{r})').font = bold
    sm.cell(r, 6, f'=IF(D{r}=0,"",D{r}/C{r}-1)').font = bold
    for rr_ in range(5, r + 1):
        for j in range(2, 6):
            sm.cell(rr_, j).number_format = money
        sm.cell(rr_, 6).number_format = pct_fmt
    for j in range(1, 8):
        sm.cell(r, j).border = Border(top=thin)
    sm.column_dimensions["A"].width = 22
    for L in "BCDEFG":
        sm.column_dimensions[L].width = 15

    # ---- Reference ---------------------------------------------------------------------------------
    rf = wb.create_sheet("Reference")
    rf["A1"], rf["A1"].font = "Reference: 2026 finish and 2027 under each view", title
    rf["A2"], rf["A2"].font = ("Sources: Toast POS net sales through Oct 14, 2026; investor budget workbook "
                               "(2027_Sales_budget_Investors.xlsx); run-rate = last 45 days seasonally adjusted; "
                               "target = each month of 2026 plus 65% of the growth in our plan."), muted
    rh = ["Store", "2025 actual (Toast)", "2026 finish: budget", "2026 finish: run-rate", "2027: budget",
          "2027: run-rate objective", "2027: target", "Target vs budget"]
    for j, h in enumerate(rh, start=1):
        cell = rf.cell(4, j, h)
        cell.font, cell.fill, cell.border = bold, head_fill, Border(bottom=thin)
        cell.alignment = Alignment(wrap_text=True, horizontal="center" if j > 1 else "left")
    r = 5
    for c in stores:
        rf.cell(r, 1, c["name"]).font = normal
        for j, k in ((2, "fy25"), (3, "b26"), (4, "rr26"), (5, "b27"), (6, "prop27")):
            cell = rf.cell(r, j, round(float(c[k]), 2))
            cell.font, cell.number_format = blue, money
        cell = rf.cell(r, 7, f"='Targets 2027'!{get_column_letter(col['Total 2027'])}{rows[c['key']]}")
        cell.font, cell.number_format = green, money
        rf.cell(r, 8, f"=G{r}-E{r}").number_format = money
        rf.cell(r, 8).font = normal
        r += 1
    rf.cell(r, 1, "Total, ten stores").font = bold
    for j in range(2, 9):
        L = get_column_letter(j)
        cell = rf.cell(r, j, f"=SUM({L}5:{L}{r - 1})")
        cell.font, cell.number_format, cell.border = bold, money, Border(top=thin)
    rf.cell(r, 1).border = Border(top=thin)
    rf.cell(r + 2, 1, "La Jolla Cove's run-rate objective and target include the budget's remodeled $125K a month "
                      "from Oct 1, 2027. New cafés are not included.").font = muted
    rf.column_dimensions["A"].width = 22
    for j in range(2, 9):
        rf.column_dimensions[get_column_letter(j)].width = 16
    rf.row_dimensions[4].height = 30

    wb.save(XLSX)
    script = next(RECALC.rglob("xlsx/scripts/recalc.py"), None)
    if script is None:
        print("note: recalc.py not found; workbook saved without cached values")
        return {"status": "not recalculated"}
    res = subprocess.run([sys.executable, str(script), str(XLSX), "60"], capture_output=True, text=True)
    try:
        out = json.loads(res.stdout[res.stdout.index("{"):])
    except ValueError:
        raise RuntimeError(f"recalc failed: {res.stdout}\n{res.stderr}")
    if out.get("status") != "success":
        raise RuntimeError(f"workbook has formula errors: {out}")
    return out


# ------------------------------------------------------------------------------------------
# Output
# ------------------------------------------------------------------------------------------
def render(n: dict, svgs: dict) -> Path:
    env = Environment(loader=FileSystemLoader(ROOT / "report"), autoescape=False)
    env.filters.update(fm=fm, fa=fa, fk=fk, sfk=sfk, sfm=sfm, pct=pct, pct0=pct0, apct=apct, frange=frange,
                       words=words)
    html = env.get_template("targets.html.j2").render(n=n, svg=svgs, BADGE=BADGE,
                                                      month_labels=[m.strftime("%b") for m in M27])
    path = BUILD / "parakeet_2027_run_rate_and_targets.html"
    path.write_text(html, encoding="utf-8")
    return path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ds", help="path to the parakeet-cafe-design skill folder")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--shots", help="also write page PNGs to this folder (visual QA)")
    args = ap.parse_args()
    ds = find_design_system(args.ds)
    charts.setup(copy_assets(ds))
    for f in bd.PHOTOS + ["photos/split-menus.jpg"]:
        shutil.copy(ds / "assets" / f, BUILD / "assets" / Path(f).name)
    n = collect()
    print("workbook:", export_xlsx(n))
    html = render(n, make_charts(n))
    print(f"HTML -> {html.relative_to(ROOT)}")
    if not args.no_pdf:
        node_path = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
        env = {**os.environ, "NODE_PATH": node_path}
        cmd = ["node", str(ROOT / "report" / "render_pdf.cjs"), str(html), str(PDF), args.shots or "",
               "13.333in", "7.5in"]
        subprocess.run(cmd, check=True, env=env)
        print(f"PDF  -> {PDF.relative_to(ROOT)}")
        print(f"XLSX -> {XLSX.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
