"""Split the 2027 monthly manager targets into days, without moving any month by a cent.

Each month's target, as sent (outputs/Parakeet_2027_Manager_Targets.xlsx, 'Targets 2027'), is spread over its days
in proportion to the café's daily seasonal pattern for 2027: the run-rate's Prophet fit (adopted settings, data
through Sep 30, 2026, src/runrate_scenario.py) for day of week, time of year and holidays, and zero on days the café
is closed. Beverly Hills' weekly pattern changed in 2026 (its weekdays caught up with its weekends, while the model,
trained mostly on earlier years, still expects strong weekends), so its days are reweighted to its actual 2026
weekday pattern. Amounts are in cents, rounded so that every month adds up exactly to its target.

    python report/build_daily_targets.py   -> outputs/Parakeet_2027_Daily_Targets.xlsx (daily targets, check against
                                              the monthly targets, weekday patterns) and outputs/targets_2027_daily.csv
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from joblib import Parallel, delayed

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from forecasting import HOLIDAYS, OUTPUTS, load_stores  # noqa: E402
from runrate_scenario import seasonal_model  # noqa: E402

MONTHLY = OUTPUTS / "Parakeet_2027_Manager_Targets.xlsx"
XLSX = OUTPUTS / "Parakeet_2027_Daily_Targets.xlsx"
CSV = OUTPUTS / "targets_2027_daily.csv"
YEAR = pd.date_range("2027-01-01", "2027-12-31")
KEY = {"Beverly Hills": "Beverly Hills", "Brentwood": "Brentwood", "Carlsbad": "Carlsbad", "Coronado": "Coronado",
       "Del Mar Café": "Carmel Valley", "La Jolla Café": "La Jolla", "La Jolla Juicery": "La Jolla Juicery",
       "Little Italy Café": "Little Italy", "Newport Beach": "Newport", "La Jolla Cove": "La Jolla Cove"}
# cafés whose weekday pattern comes from these recent full weeks (Monday to Sunday) instead of the model
WEEKLY_FROM_ACTUAL = {"Beverly Hills": ("2026-01-05", "2026-10-11")}
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
RECALC = Path.home() / ".claude" / "skills"  # the xlsx skill's LibreOffice recalculation script lives here
HOLIDAY_DATES = set(HOLIDAYS.ds)


def read_monthly() -> dict[str, np.ndarray]:
    """The monthly targets exactly as stored in the targets workbook, by café name."""
    from openpyxl import load_workbook

    ws = load_workbook(MONTHLY, data_only=True)["Targets 2027"]
    out = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] in KEY:
            out[row[0]] = np.array(row[1:13], dtype=float)
    assert set(out) == set(KEY), set(KEY) - set(out)
    return out


def weekday_index(values: pd.Series, normal: pd.Series) -> np.ndarray:
    """Average share of the week by weekday (week average = 1), over full Monday-Sunday weeks whose seven days are
    all normal trading days: geometric mean across weeks of each day's value over its week's mean."""
    d = pd.DataFrame({"v": values, "ok": normal & (values > 0)})
    d["week"] = d.index - pd.to_timedelta(d.index.dayofweek, unit="D")
    full = d.groupby("week").ok.transform(lambda x: x.all() and len(x) == 7)
    d = d[full]
    share = d.v / d.groupby("week").v.transform("mean")
    g = np.exp(np.log(share).groupby(d.index.dayofweek).mean()).reindex(range(7)).to_numpy()
    return g / g.mean()


def split_cents(total: float, weights: np.ndarray) -> np.ndarray:
    """Integer cents in proportion to the weights, adding up exactly to the total (largest remainder)."""
    cents = int(round(total * 100))
    raw = weights / weights.sum() * cents
    base = np.floor(raw).astype(np.int64)
    rest = cents - int(base.sum())
    base[np.argsort(-(raw - base), kind="stable")[:rest]] += 1
    return base


def daily_factors(sd) -> np.ndarray:
    """The café's 2027 daily seasonal factors (day of week, time of year, holidays), zero on closed days."""
    m = seasonal_model(sd)
    p = m.predict(pd.DataFrame({"ds": YEAR}))
    return (1 + p.multiplicative_terms.to_numpy()) * sd.open_flags(YEAR)


def build() -> dict:
    monthly = read_monthly()
    stores = load_stores("2026-10-14")
    names = list(monthly)
    factors = dict(zip(names, Parallel(n_jobs=4)(delayed(daily_factors)(stores[KEY[s]]) for s in names)))
    # the same seasonal pattern as the deck's run-rate: its monthly sums match the run-rate's seasonal factors
    sea = pd.read_csv(OUTPUTS / "runrate_45d_seasonal_by_month.csv").set_index(["store", "month"]).factor_sum
    months = YEAR.to_period("M")
    for s in names:
        by_month = pd.Series(factors[s], index=months).groupby(level=0).sum().to_numpy()
        assert np.allclose(by_month, [sea[(KEY[s], str(p))] for p in pd.period_range("2027-01", "2027-12", freq="M")],
                           rtol=1e-4), s

    normal_2027 = pd.Series(~YEAR.isin(HOLIDAY_DATES), index=YEAR)
    patterns, cents = {}, {}
    for s in names:
        f = factors[s]
        model_idx = weekday_index(pd.Series(f, index=YEAR), normal_2027 & (f > 0))
        patterns[s] = {"model": model_idx}
        w = f.copy()
        if s in WEEKLY_FROM_ACTUAL:
            a, b = WEEKLY_FROM_ACTUAL[s]
            d = stores[KEY[s]].daily.loc[a:b]
            normal = (d.open == 1) & ~d.anomaly & ~d.index.isin(HOLIDAY_DATES)
            actual_idx = weekday_index(d.sales.fillna(0), normal)
            patterns[s]["actual"] = actual_idx
            w = f * (actual_idx / model_idx)[YEAR.dayofweek]
        cents[s] = np.concatenate([split_cents(monthly[s][i], w[months.month == i + 1]) for i in range(12)])
        # every month (and so the year) adds up to the target as sent, to the cent
        for i in range(12):
            assert cents[s][months.month == i + 1].sum() == int(round(monthly[s][i] * 100)), (s, i)
        assert (cents[s][f == 0] == 0).all()  # nothing on days the café is closed
    return {"names": names, "monthly": monthly, "cents": cents, "patterns": patterns, "factors": factors}


def export(b: dict) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    names, monthly, cents, patterns = b["names"], b["monthly"], b["cents"], b["patterns"]
    arial = "Arial"
    title, bold, normal = (Font(name=arial, size=14, bold=True), Font(name=arial, size=10, bold=True),
                           Font(name=arial, size=10))
    blue, green = Font(name=arial, size=10, color="0000FF"), Font(name=arial, size=10, color="008000")
    muted = Font(name=arial, size=9, italic=True, color="565A5C")
    head_fill, thin = PatternFill("solid", fgColor="F2F2F2"), Side(style="thin", color="000000")
    money = '$#,##0.00;($#,##0.00);"-"'
    adjusted = ", ".join(f"{s} by its actual weekday pattern of {pd.Timestamp(a):%b %-d} – {pd.Timestamp(z):%b %-d, %Y}"
                         for s, (a, z) in WEEKLY_FROM_ACTUAL.items())

    def header(ws, row, labels):
        for j, h in enumerate(labels, start=1):
            c = ws.cell(row, j, h)
            c.font, c.fill, c.border = bold, head_fill, Border(bottom=thin)
            c.alignment = Alignment(horizontal="left" if j == 1 else "center", wrap_text=True)

    wb = Workbook()
    # ---- Daily targets ---------------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Daily targets 2027"
    ws["A1"], ws["A1"].font = "Parakeet Café · 2027 daily sales targets by store", title
    ws["A2"], ws["A2"].font = (
        "Each month's target (Parakeet_2027_Manager_Targets.xlsx) spread over its days by the café's seasonal "
        f"pattern: day of week, time of year and holidays; {adjusted}. Every month adds up exactly to its target "
        "(see 'Check'). Blue = values from the analysis; black = formulas."), muted
    header(ws, 4, ["Date", "Day", "Month"] + names + ["Total"])
    first, last = 5, 4 + len(YEAR)
    for i, day in enumerate(YEAR):
        r = first + i
        c = ws.cell(r, 1, day.to_pydatetime())
        c.number_format, c.font = "yyyy-mm-dd", normal
        ws.cell(r, 2, DAYS[day.dayofweek]).font = normal
        ws.cell(r, 3, day.strftime("%b %Y")).font = normal
        for j, s in enumerate(names):
            c = ws.cell(r, 4 + j, cents[s][i] / 100)
            c.font, c.number_format = blue, money
        tot = ws.cell(r, 4 + len(names), f"=SUM({get_column_letter(4)}{r}:{get_column_letter(3 + len(names))}{r})")
        tot.font, tot.number_format = normal, money
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 10
    for j in range(4, 5 + len(names)):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "D5"

    # ---- Check against the monthly targets as sent -------------------------------------------------------
    ck = wb.create_sheet("Check")
    ck["A1"], ck["A1"].font = "Check: the daily targets add up to the monthly targets as sent", title
    ck["A2"], ck["A2"].font = ("Monthly target: 'Targets 2027' in Parakeet_2027_Manager_Targets.xlsx (blue). "
                               "Sum of daily: this workbook (green). Difference in cents: must be 0."), muted
    ck["A3"], ck["A3"].font = "Total absolute difference (cents)", bold
    header(ck, 5, ["Store", "Month", "Monthly target (as sent)", "Sum of daily targets", "Difference (cents)"])
    r = 6
    for j, s in enumerate(names):
        col = get_column_letter(4 + j)
        start = r
        for i, label in enumerate(pd.period_range("2027-01", "2027-12", freq="M").strftime("%b %Y")):
            ck.cell(r, 1, s).font = normal
            ck.cell(r, 2, label).font = normal
            c = ck.cell(r, 3, float(monthly[s][i]))
            c.font, c.number_format = blue, money
            c = ck.cell(r, 4, f"=SUMIFS('Daily targets 2027'!{col}${first}:{col}${last},"
                              f"'Daily targets 2027'!$C${first}:$C${last},B{r})")
            c.font, c.number_format = green, money
            ck.cell(r, 5, f"=ROUND(D{r}*100,0)-ROUND(C{r}*100,0)").font = normal
            r += 1
        ck.cell(r, 1, s).font = bold
        ck.cell(r, 2, "Year").font = bold
        for k, L in ((3, "C"), (4, "D")):
            c = ck.cell(r, k, f"=SUM({L}{start}:{L}{r - 1})")
            c.font, c.number_format = bold, money
        ck.cell(r, 5, f"=ROUND(D{r}*100,0)-ROUND(C{r}*100,0)").font = bold
        for k in range(1, 6):
            ck.cell(r, k).border = Border(top=thin)
        r += 2
    ck["B3"] = f"=SUMPRODUCT(ABS(E6:E{r - 2}))"
    ck["B3"].font = bold
    ck.column_dimensions["A"].width = 30
    for L, wdt in (("B", 12), ("C", 22), ("D", 22), ("E", 18)):
        ck.column_dimensions[L].width = wdt
    ck.freeze_panes = "A6"

    # ---- Weekday patterns ----------------------------------------------------------------------------------
    wp = wb.create_sheet("Weekday pattern")
    wp["A1"], wp["A1"].font = "Weekday pattern behind the daily split (week average = 100)", title
    wp["A2"], wp["A2"].font = ("Model: the café's 2027 seasonal pattern from the run-rate's Prophet fit (full weeks "
                               "without holidays). Used for the daily targets unless a row says otherwise."), muted
    header(wp, 4, ["Store", "Pattern"] + DAYS)
    r = 5
    for s in names:
        rows = [("Model", patterns[s]["model"])]
        if "actual" in patterns[s]:
            a, z = WEEKLY_FROM_ACTUAL[s]
            rows = [("Model (not used)", patterns[s]["model"]),
                    (f"Used: actual {pd.Timestamp(a):%b %-d} – {pd.Timestamp(z):%b %-d, %Y}", patterns[s]["actual"])]
        for label, v in rows:
            wp.cell(r, 1, s).font = normal
            wp.cell(r, 2, label).font = normal
            for k in range(7):
                c = wp.cell(r, 3 + k, round(float(v[k]) * 100, 1))
                c.font, c.number_format = blue, "0.0"
            r += 1
    wp.cell(r + 1, 1, "Beverly Hills: its weekdays caught up with its weekends in 2026 (Tuesday to Sunday all about "
                      "100, against a Saturday of about 116 in 2025), and the model still expects strong weekends. "
                      "The monthly targets are unchanged; only the split across days moves.").font = muted
    wp.column_dimensions["A"].width = 20
    wp.column_dimensions["B"].width = 30
    for k in range(7):
        wp.column_dimensions[get_column_letter(3 + k)].width = 8

    wb.save(XLSX)
    script = next(RECALC.rglob("xlsx/scripts/recalc.py"), None)
    if script is None:
        print("note: recalc.py not found; workbook saved without cached values")
        return {"status": "not recalculated"}
    res = subprocess.run([sys.executable, str(script), str(XLSX), "120"], capture_output=True, text=True)
    try:
        out = json.loads(res.stdout[res.stdout.index("{"):])
    except ValueError:
        raise RuntimeError(f"recalc failed: {res.stdout}\n{res.stderr}")
    if out.get("status") != "success":
        raise RuntimeError(f"workbook has formula errors: {out}")
    return out


def verify_workbook() -> float:
    """Read the recalculated check: the total absolute difference in cents must be zero."""
    from openpyxl import load_workbook

    return load_workbook(XLSX, data_only=True)["Check"]["B3"].value


def main() -> None:
    b = build()
    rows = [{"date": d.date(), "weekday": DAYS[d.dayofweek], "store": s, "target": b["cents"][s][i] / 100}
            for s in b["names"] for i, d in enumerate(YEAR)]
    pd.DataFrame(rows).to_csv(CSV, index=False, float_format="%.2f")
    print("workbook:", export(b))
    diff = verify_workbook()
    assert diff == 0, diff
    print(f"check: total absolute difference {diff} cents")
    for s in b["names"]:
        print(f"{s:18s} {sum(b['cents'][s]) / 100:>14,.2f}  (as sent {b['monthly'][s].sum():>14,.2f})")
    print(f"-> {XLSX.relative_to(ROOT)}\n-> {CSV.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
